from pathlib import Path
import os

from docx import Document


DOCX = Path(
    "/Users/nguyenhaiduong/Downloads/Student-system/"
    "PTTK_Student_Digital_Competency_Profile_Ban_Chuyen_Nghiep_Toan_He_Thong.docx"
)

PAGE_BY_CHAPTER = {
    "1. Giới thiệu, bối cảnh và mục tiêu": "4",
    "2. Phạm vi, tác nhân và giả định": "5",
    "3. Phân tích nghiệp vụ và yêu cầu": "6",
    "4. Phân rã chức năng cấp 1/2/3": "10",
    "5. Use Case toàn hệ thống": "19",
    "6. Kiến trúc logic": "37",
    "7. Thiết kế xử lý và luồng dữ liệu": "39",
    "8. Thiết kế dữ liệu toàn hệ thống": "42",
    "9. Thiết kế sâu module nghiệp vụ": "81",
    "10. Competency Analysis Engine": "85",
    "11. Hợp đồng API thiết kế đích": "94",
    "12. Thiết kế giao diện theo vai trò": "109",
    "13. Phân quyền và bảo mật": "114",
    "14. Kiểm thử và nghiệm thu": "117",
    "15. Ma trận truy vết yêu cầu": "121",
    "16. Phụ lục": "123",
}


def replace_cell_text(cell, value: str):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def main():
    document = Document(DOCX)
    updated = 0
    for table in document.tables:
        if not table.rows or len(table.rows[0].cells) < 2:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:2] != ["Nội dung", "Trang"]:
            continue
        for row in table.rows[1:]:
            title = row.cells[0].text.strip()
            if title in PAGE_BY_CHAPTER:
                replace_cell_text(row.cells[1], PAGE_BY_CHAPTER[title])
                updated += 1
        break

    if updated != len(PAGE_BY_CHAPTER):
        raise RuntimeError(
            f"Chỉ cập nhật {updated}/{len(PAGE_BY_CHAPTER)} dòng mục lục."
        )

    tmp = DOCX.with_suffix(".toc-pages.tmp.docx")
    document.save(tmp)
    os.replace(tmp, DOCX)
    print(f"Đã cập nhật {updated} số trang trong mục lục.")


if __name__ == "__main__":
    main()
