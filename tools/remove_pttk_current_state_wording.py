#!/usr/bin/env python3
"""Loại bỏ dấu vết ghép tài liệu và chuẩn hóa nội dung thành thiết kế đích."""

from __future__ import annotations

import os
from pathlib import Path

from docx import Document


DOCX = Path(
    "/Users/nguyenhaiduong/Downloads/Student-system/"
    "PTTK_Student_Digital_Competency_Profile_Ban_Chuyen_Nghiep_Toan_He_Thong.docx"
)


PARAGRAPH_REPLACEMENTS = {
    "Từ điển dưới đây được đối chiếu với schema Prisma hiện tại. Thiết kế đích giữ ràng buộc hiện có và đánh dấu các bảng đề xuất bổ sung ở mục 18.4.":
        "Từ điển dữ liệu dưới đây mô tả thiết kế vật lý đích, gồm kiểu dữ liệu, ràng buộc, chỉ mục và ý nghĩa nghiệp vụ của từng trường.",
}

CELL_REPLACEMENTS = {
    "Loại bỏ hằng số khỏi service": "Tập trung cấu hình, tránh hằng số rải rác.",
    "Đã có trong schema và service.": "Thuộc phạm vi thiết kế phiên bản đầu.",
    "Cần bổ sung API/UI lịch sử.": "API/UI lịch sử phải hiển thị đầy đủ các lần làm.",
    "Các hằng số 6 câu/skill, 60%, 3 lần, 30 phút và 60 phút hiện nằm trong service. Thiết kế đích nên đưa chúng vào AssessmentPolicy có version để lịch sử bài luôn giải thích được.":
        "Các tham số 6 câu/skill, 60%, 3 lần, 30 phút và 60 phút phải được cấu hình trong AssessmentPolicy có phiên bản để lịch sử bài luôn giải thích được.",
    "Đề xuất thêm RecommendationFeedback.": "Lưu phản hồi trong RecommendationFeedback.",
    "Đề xuất bổ sung để chặn task học sai thứ tự.": "Bắt buộc áp dụng để chặn task học sai thứ tự.",
    "Đề xuất thêm ExportJob thay vì chỉ ExportLog": "Thực hiện qua ExportJob; ExportLog lưu dấu vết.",
}

DECISIONS = [
    "AssessmentPolicy và AnalysisPolicy được lưu trong cơ sở dữ liệu theo phiên bản; bản phát hành chỉ cung cấp cấu trúc và dữ liệu khởi tạo mặc định.",
    "Tài khoản sinh viên, cố vấn và quản trị do nhà trường tạo hoặc nhập theo lô; không mở đăng ký công khai trong phiên bản đầu.",
    "Thời hạn minh chứng được cấu hình theo loại; cố vấn được phân công ACTIVE hoặc quản trị viên xác minh trong tối đa 03 ngày làm việc.",
    "Mức tự khai chỉ là dữ liệu tham khảo với sourceConfidence mặc định 0,25; không tự động ghi đè currentLevel từ assessment hoặc minh chứng đã xác minh.",
    "Lịch sử assessment hiển thị điểm, mức đạt và lời giải sau khi bài được chấm; đáp án không được trả về khi bài còn IN_PROGRESS.",
    "Giới hạn số lần đánh giá được tính theo student, specialization và policyVersion; phiên bản chính sách mới bắt đầu một chu kỳ giới hạn mới.",
    "Phản hồi gợi ý chỉ phục vụ báo cáo chất lượng trong phiên bản đầu; mọi thay đổi ranking hoặc template phải do quản trị viên phê duyệt và tạo phiên bản mới.",
    "Dependency HARD chặn bắt đầu RoadmapItem khi mục tiền nhiệm chưa DONE; dependency SOFT chỉ cảnh báo và cho phép tiếp tục khi người dùng xác nhận.",
    "Portfolio chỉ công khai các trường do sinh viên chọn; liên kết và mã QR dùng URL ký có thời hạn 07 ngày, không tạo hồ sơ công khai vĩnh viễn.",
    "Báo cáo XLSX giới hạn 50.000 dòng mỗi tệp; dữ liệu lớn hơn được chia tệp và đóng gói ZIP trong cùng ExportJob.",
    "ActivityLog và EvidenceRevision được lưu tối thiểu 05 năm; tệp export lưu 07 ngày, còn metadata ExportJob/ExportLog lưu 01 năm.",
    "Kênh IN_APP là bắt buộc trong phiên bản đầu; EMAIL được bật theo cấu hình cho thông báo ưu tiên cao và không thay thế bản ghi Notification.",
]


def remove_element(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def replace_text(paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    paragraph.text = paragraph.text.replace(old, new)
    return True


def main() -> None:
    document = Document(DOCX)

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text in PARAGRAPH_REPLACEMENTS:
            paragraph.text = PARAGRAPH_REPLACEMENTS[text]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for old, new in CELL_REPLACEMENTS.items():
                        replace_text(paragraph, old, new)

    duplicate = (
        "Ma trận nối yêu cầu với use case, business rule, API, bảng dữ liệu, "
        "màn hình và test. Cột ghi chú chỉ mô tả quyết định thiết kế, không "
        "phản ánh trạng thái triển khai."
    )
    duplicate_paragraphs = [p for p in document.paragraphs if p.text.strip() == duplicate]
    for paragraph in duplicate_paragraphs[1:]:
        remove_element(paragraph._element)

    # Xóa trang phân cách “Phần II” và các hộp mô tả từng dùng để ghép phần bổ sung.
    cover_texts = {
        "PHẦN II",
        "ĐẶC TẢ CHI TIẾT TOÀN HỆ THỐNG",
        "Từ chức năng → Use Case → API → Dữ liệu → Giao diện → Kiểm thử → Mã nguồn",
    }
    cover_paragraphs = []
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() in cover_texts:
            cover_paragraphs.extend(paragraphs[max(0, index - 3): min(len(paragraphs), index + 5)])
    for paragraph in list(dict.fromkeys(cover_paragraphs)):
        if paragraph.text.strip() in cover_texts or not paragraph.text.strip():
            remove_element(paragraph._element)

    for table in list(document.tables):
        text = "\n".join(cell.text.strip() for row in table.rows for cell in row.cells)
        if text.startswith("Kết thúc Phần I") or text.startswith("Mục tiêu nâng cấp"):
            remove_element(table._element)

    heading = next(
        p for p in document.paragraphs
        if p.text.strip() == "16.4. Câu hỏi cần chốt trước triển khai"
    )
    heading.text = "16.4. Các quyết định triển khai mặc định"
    numbered = []
    started = False
    for paragraph in document.paragraphs:
        if paragraph._p is heading._p:
            started = True
            continue
        if not started:
            continue
        if paragraph.style.name == "List Number":
            numbered.append(paragraph)
            if len(numbered) == len(DECISIONS):
                break
    if len(numbered) != len(DECISIONS):
        raise RuntimeError(f"Không tìm đủ danh sách quyết định: {len(numbered)}")
    for paragraph, decision in zip(numbered, DECISIONS):
        paragraph.text = decision

    temp = DOCX.with_suffix(".tmp.docx")
    document.save(temp)
    os.replace(temp, DOCX)
    print("Đã chuẩn hóa nội dung thiết kế đích và xóa dấu vết ghép tài liệu.")


if __name__ == "__main__":
    main()
