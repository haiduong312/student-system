from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph

from refine_pttk_target_design import (
    DOCX,
    ENTITY_META,
    TMP,
    Writer,
    add_entity_fields,
    add_plain_table,
    add_traceability,
    format_table,
    renumber_captions,
    renumber_headings,
    sanitize_text,
    update_front_matter,
    validate_design_only,
)


ROOT = Path(__file__).resolve().parents[1]
MEDIA = ROOT / ".pttk_media_recover/word/media"


def find_heading(doc: Document, startswith: str):
    for p in doc.paragraphs:
        if p.style and p.style.name.startswith("Heading") and p.text.strip().startswith(startswith):
            return p
    raise RuntimeError(f"Không tìm thấy heading: {startswith}")


def move_new_elements_before(doc: Document, start_count: int, target_el) -> None:
    body = doc.element.body
    new_elements = list(body)[start_count:-1]
    for el in new_elements:
        target_el.addprevious(el)


def insert_context_for(doc: Document, entity_name: str) -> None:
    target_heading = None
    target_table = None
    active = False
    for el in list(doc.element.body):
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            if p.style and p.style.name == "Heading 3" and p.text.strip().endswith(f"Bảng {entity_name}"):
                target_heading = el
                active = True
        elif el.tag == qn("w:tbl") and active:
            table = Table(el, doc)
            if table.rows and table.cell(0, 0).text.strip() in {"Field", "Trường"}:
                target_table = el
                break
    if target_heading is None or target_table is None:
        raise RuntimeError(f"Không tìm thấy bảng dữ liệu {entity_name}")
    purpose, relations, states, unique_index, use_cases, apis, tests = ENTITY_META[entity_name]
    caption = doc.add_paragraph(f"Bảng 0. Ngữ cảnh nghiệp vụ {entity_name}", style="Caption")
    table = add_plain_table(
        doc,
        ["Thuộc tính thiết kế", "Đặc tả"],
        [["Mục đích nghiệp vụ", purpose], ["Quan hệ chính", relations], ["Trạng thái và ý nghĩa", states],
         ["Unique/index", unique_index], ["Use case tạo/sửa/xóa", use_cases], ["API liên quan", apis],
         ["Test case liên quan", tests]],
        [2050, 7310],
        7.8,
    )
    target_heading.addnext(table._tbl)
    target_heading.addnext(caption._p)


def add_missing_data_and_trace(doc: Document) -> None:
    chapter9 = find_heading(doc, "9.")
    start = len(list(doc.element.body)) - 1
    w = Writer(doc)
    add_entity_fields(w, "RoadmapItemDependency", [
        ["roadmapItemId", "UUID", "Có", "FK RoadmapItem", "Mục bị phụ thuộc."],
        ["predecessorItemId", "UUID", "Có", "FK RoadmapItem", "Mục phải hoàn thành trước."],
        ["dependencyType", "Enum", "Có", "HARD/SOFT", "HARD chặn bắt đầu; SOFT chỉ cảnh báo."],
        ["createdAt", "DateTime", "Có", "-", "Thời điểm tạo quan hệ."],
        ["(constraint)", "-", "Có", "unique item+predecessor; check khác nhau", "Không tạo vòng lặp hoặc tự phụ thuộc."],
    ])
    add_entity_fields(w, "ExportJob", [
        ["id", "UUID", "Có", "PK", "Định danh job."], ["ownerId", "UUID", "Có", "FK User; index", "Chủ sở hữu kết quả xuất."],
        ["exportType", "Enum", "Có", "index", "PORTFOLIO_PDF hoặc REPORT_XLSX."], ["filters", "JSON", "Có", "snapshot", "Bộ lọc đã áp dụng scope."],
        ["status", "ExportJobStatus", "Có", "index", "QUEUED/PROCESSING/COMPLETED/FAILED/EXPIRED."], ["progress", "Int", "Có", "0..100", "Phần trăm xử lý."],
        ["fileUrl", "String?", "Không", "signed/private", "Đường dẫn tải có hạn."], ["checksum", "String?", "Không", "-", "Kiểm tra toàn vẹn tệp."],
        ["idempotencyKey", "String", "Có", "unique ownerId+key", "Chống tạo job trùng."], ["errorCode", "String?", "Không", "index", "Mã lỗi ổn định cho client."],
        ["expiresAt", "DateTime?", "Không", "index", "Thời điểm thu hồi file."], ["createdAt", "DateTime", "Có", "index", "Thời điểm tạo job."],
    ])
    move_new_elements_before(doc, start, chapter9._p)
    insert_context_for(doc, "RoadmapItemDependency")
    insert_context_for(doc, "ExportJob")

    chapter16 = find_heading(doc, "16.")
    start = len(list(doc.element.body)) - 1
    w = Writer(doc)
    add_traceability(w)
    move_new_elements_before(doc, start, chapter16._p)


def insert_heading_before_table(doc: Document, expected_headers: list[str], heading_text: str) -> None:
    for el in list(doc.element.body):
        if el.tag != qn("w:tbl"):
            continue
        table = Table(el, doc)
        headers = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
        if headers[:len(expected_headers)] == expected_headers:
            previous = el.getprevious()
            while previous is not None and previous.tag == qn("w:p"):
                p = Paragraph(previous, doc)
                if p.text.strip():
                    break
                previous = previous.getprevious()
            if previous is not None and previous.tag == qn("w:p"):
                p = Paragraph(previous, doc)
                if p.style and p.style.name == "Caption":
                    target = previous
                else:
                    target = el
            else:
                target = el
            new_p = doc.add_paragraph(heading_text, style="Heading 2")
            target.addprevious(new_p._p)
            return


def cleanup_structure(doc: Document) -> None:
    for p in doc.paragraphs:
        if p.text.strip().endswith("Các bảng đề xuất bổ sung"):
            p.text = "Các bảng chính sách và vòng đời dữ liệu"
        elif p.style and p.style.name == "Heading 2" and p.text.strip().endswith("Chiến lược kiểm thử"):
            if p.text.strip().startswith("14.3"):
                p.text = "Các lớp kiểm thử và dữ liệu thử"

    # Xóa ma trận rút gọn ở Chương 14 vì Chương 15 là ma trận đầy đủ.
    start = None
    end = None
    for el in list(doc.element.body):
        if el.tag != qn("w:p"):
            continue
        p = Paragraph(el, doc)
        if p.style and p.style.name == "Heading 2" and "Ma trận truy vết rút gọn" in p.text:
            start = el
        elif start is not None and p.style and p.style.name == "Heading 1" and p.text.strip().startswith("15."):
            end = el
            break
    if start is not None and end is not None:
        cur = start
        while cur is not end:
            nxt = cur.getnext()
            cur.getparent().remove(cur)
            cur = nxt

    insert_heading_before_table(doc, ["Mã", "Nhóm", "Tiêu chí đo"], "Yêu cầu phi chức năng")


def add_picture_after(doc: Document, heading_prefix: str, image_name: str, caption_text: str, alt_text: str) -> None:
    heading = find_heading(doc, heading_prefix)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    shape = run.add_picture(str(MEDIA / image_name), width=Inches(6.25))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    caption = doc.add_paragraph(f"Hình 0. {caption_text}", style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading._p.addnext(caption._p)
    heading._p.addnext(p._p)


def restore_figures(doc: Document) -> None:
    mappings = [
        ("1.1.", "image1.png", "Bối cảnh nghiệp vụ của hệ thống", "Sơ đồ tác nhân và đầu ra nghiệp vụ của hệ thống."),
        ("5.1.", "image7.png", "Mô hình Use Case toàn hệ thống", "Sơ đồ use case với tác nhân Sinh viên, Cố vấn và Quản trị viên."),
        ("6.1.", "image2.png", "Kiến trúc logic đề xuất", "Sơ đồ các lớp trình bày, ứng dụng, miền nghiệp vụ và hạ tầng."),
        ("7.1.", "image4.png", "Sơ đồ luồng dữ liệu mức đỉnh", "Sơ đồ tác nhân, tiến trình P1-P4 và các kho dữ liệu D1-D4."),
        ("7.3.", "image5.png", "Trình tự thực hiện phân tích năng lực", "Sơ đồ trình tự giữa Sinh viên, dịch vụ ứng dụng, Engine và kho dữ liệu."),
        ("8.1.", "image6.png", "Mô hình dữ liệu khái niệm theo cụm miền", "Sơ đồ các cụm danh tính, danh mục, đánh giá, phân tích và phát triển cá nhân."),
        ("8.6.", "image8.png", "Bản đồ dữ liệu toàn hệ thống", "Bản đồ các miền dữ liệu và luồng phụ thuộc chính."),
        ("10.6.", "image3.png", "Chuỗi xử lý của Competency Analysis Engine", "Sơ đồ bảy bước xử lý từ ngữ cảnh đến lưu snapshot kết quả."),
    ]
    for args in mappings:
        add_picture_after(doc, *args)


def main() -> None:
    doc = Document(DOCX)
    add_missing_data_and_trace(doc)
    cleanup_structure(doc)
    restore_figures(doc)
    sanitize_text(doc)
    renumber_headings(doc)
    renumber_captions(doc)
    update_front_matter(doc)
    validate_design_only(doc)
    doc.save(TMP)
    os.replace(TMP, DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
