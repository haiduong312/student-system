from pathlib import Path
import os
import re

from docx import Document


DOCX = Path(
    "/Users/nguyenhaiduong/Downloads/Student-system/"
    "PTTK_Student_Digital_Competency_Profile_Ban_Chuyen_Nghiep_Toan_He_Thong.docx"
)


GENERIC_CAPTIONS = {
    "Bối cảnh nghiệp vụ của hệ thống",
    "Mô hình Use Case toàn hệ thống",
    "Kiến trúc logic đề xuất",
    "Sơ đồ luồng dữ liệu mức đỉnh",
    "Trình tự thực hiện phân tích năng lực",
    "Mô hình dữ liệu khái niệm theo cụm miền",
    "Bản đồ dữ liệu toàn hệ thống",
    "Chuỗi xử lý của Competency Analysis Engine",
}


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def caption_body(text: str) -> str | None:
    match = re.match(r"^Hình\s+\d+\s*[.:–-]?\s*(.+?)\s*$", text)
    return match.group(1) if match else None


def renumber_figure_captions(document: Document) -> int:
    number = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        body = caption_body(text)
        if body is None:
            continue
        number += 1
        paragraph.text = f"Hình {number}. {body}"
        paragraph.style = "Caption"
    return number


def main():
    document = Document(DOCX)
    removed = 0
    for paragraph in list(document.paragraphs):
        body = caption_body(paragraph.text.strip())
        if body in GENERIC_CAPTIONS:
            remove_paragraph(paragraph)
            removed += 1

    remaining = renumber_figure_captions(document)
    tmp = DOCX.with_suffix(".caption-fix.tmp.docx")
    document.save(tmp)
    os.replace(tmp, DOCX)
    print(f"Đã xóa {removed} caption lặp; còn {remaining} caption hình.")


if __name__ == "__main__":
    main()
