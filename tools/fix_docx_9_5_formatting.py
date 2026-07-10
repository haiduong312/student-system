from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


TABLE_WIDTH = 9360
TABLE_INDENT = 120
MID_GRAY = "D2DAE5"
HEADER_FILL = "F2F4F7"


TABLE_WIDTHS: dict[tuple[str, ...], list[int]] = {
    ("Nhóm dữ liệu", "Thành phần", "Vai trò trong sinh lộ trình"): [1900, 4400, 3060],
    ("Mã điều kiện", "Điều kiện", "Xử lý nếu không đạt"): [1450, 3850, 4060],
    ("Skill", "CurrentLevel", "RequiredLevel", "Gap", "Weight", "PriorityScore", "Kết luận"): [850, 1200, 1350, 700, 850, 1350, 3060],
    ("RecommendationType", "RoadmapItemType", "Ý nghĩa"): [3000, 2800, 3560],
    ("Phase", "Ý nghĩa", "Thứ tự"): [3200, 4600, 1560],
    ("Tình huống", "Quy tắc xử lý", "Thông báo gợi ý"): [2200, 2700, 4460],
    ("Trạng thái item", "Ý nghĩa", "Ảnh hưởng đến tiến độ"): [2000, 2100, 5260],
    ("Skill", "Current", "Required", "Gap", "Weight", "PriorityScore", "Phase"): [900, 900, 1000, 700, 800, 1260, 3800],
    ("Order", "Phase", "RoadmapItem", "Skill", "Type", "Dependency", "Giải thích"): [750, 1400, 2050, 800, 1000, 1450, 1910],
    ("Mã test", "Mục tiêu kiểm thử", "Dữ liệu đầu vào", "Kết quả mong đợi"): [1100, 2200, 3000, 3060],
}


def copy_ppr(source, target) -> None:
    if target._p.pPr is not None:
        target._p.remove(target._p.pPr)
    if source._p.pPr is not None:
        target._p.insert(0, deepcopy(source._p.pPr))


def copy_rpr(source, target) -> None:
    if target._r.rPr is not None:
        target._r.remove(target._r.rPr)
    if source._r.rPr is not None:
        target._r.insert(0, deepcopy(source._r.rPr))


def set_run_size(run, half_points: int) -> None:
    r_pr = run._r.get_or_add_rPr()
    for tag in ("sz", "szCs"):
        node = r_pr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            r_pr.append(node)
        node.set(qn("w:val"), str(half_points))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str | None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if fill is None:
        if shd is not None:
            tc_pr.remove(shd)
        return
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        tr_pr.append(node)
    node.set(qn("w:val"), "true")


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), MID_GRAY)


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH:
        raise ValueError(f"Invalid table width total: {sum(widths)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def format_table(table, widths, header_p, header_r, body_p_left, body_p_center, body_r) -> None:
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    dense_table_size = 18 if len(widths) >= 7 else None

    for row_index, row in enumerate(table.rows):
        set_cant_split(row)
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, HEADER_FILL if row_index == 0 else None)
            tc_pr = cell._tc.get_or_add_tcPr()
            hide_mark = tc_pr.find(qn("w:hideMark"))
            if hide_mark is not None:
                tc_pr.remove(hide_mark)

            for paragraph in cell.paragraphs:
                if row_index == 0:
                    copy_ppr(header_p, paragraph)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        copy_rpr(header_r, run)
                        if dense_table_size is not None:
                            set_run_size(run, dense_table_size)
                else:
                    center_first = col_index == 0 and len(cell.text.strip()) <= 18
                    copy_ppr(body_p_center if center_first else body_p_left, paragraph)
                    for run in paragraph.runs:
                        copy_rpr(body_r, run)
                        if dense_table_size is not None:
                            set_run_size(run, dense_table_size)


def add_space_after_formatted_tables(doc) -> None:
    target_tbls = {
        table._tbl
        for table in doc.tables
        if table.rows and tuple(cell.text.strip() for cell in table.rows[0].cells) in TABLE_WIDTHS
    }
    body_children = list(doc.element.body)
    for index, child in enumerate(body_children[:-1]):
        if child not in target_tbls:
            continue
        following = body_children[index + 1]
        if following.tag != qn("w:p"):
            continue
        p_pr = following.get_or_add_pPr()
        spacing = p_pr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            p_pr.append(spacing)
        spacing.set(qn("w:before"), "120")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    doc = Document(str(args.input))
    text_before = tuple(node.text or "" for node in doc.element.xpath(".//w:t"))
    paragraph_count_before = len(doc.paragraphs)
    table_shapes_before = tuple((len(t.rows), len(t.columns)) for t in doc.tables)

    start = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("9.5.1."))
    end = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("9.6."))

    heading_template = next(p for p in doc.paragraphs if p.text.strip().startswith("10.6.1."))
    heading_run_template = next(r for r in heading_template.runs if r.text)
    normal_template = next(p for p in doc.paragraphs if p.text.strip().startswith("Một kỹ năng có thể có nhiều quan sát."))
    normal_run_template = next(r for r in normal_template.runs if r.text)
    list_template = next(p for p in doc.paragraphs if p.text.strip().startswith("Chuẩn hóa danh mục kỹ năng số"))
    list_run_template = next(r for r in list_template.runs if r.text)

    heading_count = 0
    for paragraph in doc.paragraphs[start:end]:
        text = paragraph.text.strip()
        if text.startswith("9.5."):
            copy_ppr(heading_template, paragraph)
            for run in paragraph.runs:
                copy_rpr(heading_run_template, run)
            heading_count += 1
            continue
        if not text:
            continue
        if any((run.font.name or "") == "Courier New" for run in paragraph.runs if run.text):
            continue

        num_pr = paragraph._p.pPr.find(qn("w:numPr")) if paragraph._p.pPr is not None else None
        saved_num_pr = deepcopy(num_pr) if num_pr is not None else None
        if saved_num_pr is not None:
            copy_ppr(list_template, paragraph)
            current_num_pr = paragraph._p.pPr.find(qn("w:numPr"))
            paragraph._p.pPr.replace(current_num_pr, saved_num_pr)
            for run in paragraph.runs:
                copy_rpr(list_run_template, run)
        else:
            copy_ppr(normal_template, paragraph)
            for run in paragraph.runs:
                copy_rpr(normal_run_template, run)

    if heading_count != 12:
        raise RuntimeError(f"Expected 12 headings, found {heading_count}")

    standard_table = next(
        t for t in doc.tables
        if tuple(c.text.strip() for c in t.rows[0].cells) == ("Vai trò", "Chỉ số", "Phạm vi", "Hành động")
    )
    header_p = standard_table.cell(0, 0).paragraphs[0]
    header_r = next(r for r in header_p.runs if r.text)
    body_p_center = standard_table.cell(1, 0).paragraphs[0]
    body_p_left = standard_table.cell(1, 1).paragraphs[0]
    body_r = next(r for r in body_p_left.runs if r.text)

    formatted_tables = 0
    for table in doc.tables:
        if not table.rows:
            continue
        headers = tuple(cell.text.strip() for cell in table.rows[0].cells)
        widths = TABLE_WIDTHS.get(headers)
        if widths is None:
            continue
        format_table(table, widths, header_p, header_r, body_p_left, body_p_center, body_r)
        formatted_tables += 1

    if formatted_tables != len(TABLE_WIDTHS):
        raise RuntimeError(f"Expected {len(TABLE_WIDTHS)} tables, formatted {formatted_tables}")
    add_space_after_formatted_tables(doc)

    text_after = tuple(node.text or "" for node in doc.element.xpath(".//w:t"))
    if text_after != text_before:
        raise RuntimeError("Document text changed during formatting")
    if len(doc.paragraphs) != paragraph_count_before:
        raise RuntimeError("Paragraph count changed during formatting")
    if tuple((len(t.rows), len(t.columns)) for t in doc.tables) != table_shapes_before:
        raise RuntimeError("Table structure changed during formatting")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(args.output))
    print(f"Saved: {args.output}")
    print(f"Formatted headings: {heading_count}")
    print(f"Formatted tables: {formatted_tables}")
    print(f"Text nodes preserved: {len(text_before)}")


if __name__ == "__main__":
    main()
