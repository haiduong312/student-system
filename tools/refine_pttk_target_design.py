from __future__ import annotations

import os
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from upgrade_pttk_full_system import (
    ASSESSMENT_POLICY,
    FUNCTION_TREE,
    RBAC_ROWS,
    TEST_CASES,
    TRACEABILITY,
    UC_DETAILS,
    Writer,
    LIGHT_BLUE,
    LIGHT_GRAY,
    GREEN,
    TABLE_WIDTH,
    set_cant_split,
    set_cell_margins,
    set_repeat_header,
    set_run_font,
    set_table_borders,
    set_table_geometry,
    shade_cell,
)


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "PTTK_Student_Digital_Competency_Profile_Ban_Chuyen_Nghiep_Toan_He_Thong.docx"
TMP = ROOT / f".{DOCX.stem}.tmp.docx"


CHAPTERS = [
    "Giới thiệu, bối cảnh và mục tiêu",
    "Phạm vi, tác nhân và giả định",
    "Phân tích nghiệp vụ và yêu cầu",
    "Phân rã chức năng cấp 1/2/3",
    "Use Case toàn hệ thống",
    "Kiến trúc logic",
    "Thiết kế xử lý và luồng dữ liệu",
    "Thiết kế dữ liệu toàn hệ thống",
    "Thiết kế sâu module nghiệp vụ",
    "Competency Analysis Engine",
    "Hợp đồng API thiết kế đích",
    "Thiết kế giao diện theo vai trò",
    "Phân quyền và bảo mật",
    "Kiểm thử và nghiệm thu",
    "Ma trận truy vết yêu cầu",
    "Phụ lục",
]


GROUP_META = {
    "F1": ("Người dùng/Quản trị viên", "Thông tin xác thực, vai trò hoặc đối tượng được phân công", "Phiên hoặc bản ghi quyền hợp lệ", "BR-ACC-01..08"),
    "F2": ("Quản trị viên", "Dữ liệu danh mục và cấu hình nghề nghiệp", "Danh mục có phiên bản và trạng thái rõ ràng", "BR-CAT-01..10"),
    "F3": ("Sinh viên/Cố vấn", "Dữ liệu hồ sơ và tệp hoặc URL minh chứng", "Hồ sơ cập nhật, trạng thái xác minh được tính lại", "BR-PRO-01..08; BR-EVD-01..10"),
    "F4": ("Sinh viên/Quản trị viên/Hệ thống", "Câu hỏi, chính sách và câu trả lời", "Bài đánh giá, kết quả theo kỹ năng và lịch sử", "BR-ASM-01..16"),
    "F5": ("Sinh viên/Cố vấn/Hệ thống", "Context, mức hiện tại, mức yêu cầu và policy", "Kết quả phân tích, chi tiết và giải thích", "BR-ENG-01..12"),
    "F6": ("Sinh viên/Quản trị viên/Hệ thống", "Khoảng cách kỹ năng, mẫu gợi ý và tài nguyên", "Gợi ý cá nhân và tiến độ học tập", "BR-REC-01..12"),
    "F7": ("Sinh viên/Hệ thống", "Kết quả phân tích và gợi ý có thể hành động", "Lộ trình, mục lộ trình và tiến độ", "BR-ROA-01..12"),
    "F8": ("Cố vấn", "Phân công ACTIVE và dữ liệu sinh viên", "Nhận xét và thông báo cho sinh viên", "BR-ADV-01..06"),
    "F9": ("Sinh viên/Cố vấn/Quản trị viên", "Dữ liệu tổng hợp đúng scope", "Chỉ số, cảnh báo và liên kết hành động", "BR-DSH-01..08"),
    "F10": ("Người dùng/Hệ thống", "Sự kiện miền và tùy chọn nhận thông báo", "Thông báo không trùng và trạng thái đọc", "BR-NOT-01..08"),
    "F11": ("Sinh viên/Cố vấn/Quản trị viên", "Phạm vi, bộ lọc, cột và định dạng", "ExportJob và tệp có thời hạn", "BR-EXP-01..10"),
    "F12": ("Quản trị viên/Hệ thống", "Actor, action, entity, requestId và thời gian", "Nhật ký bất biến và cảnh báo vận hành", "BR-AUD-01..08"),
}


ENTITY_META = {
    "User": ("Danh tính đăng nhập và vòng đời tài khoản.", "1-1 Student hoặc Lecturer; 1-n ActivityLog/Notification.", "ACTIVE cho phép tạo phiên; INACTIVE chặn đăng nhập và làm mới phiên.", "email và mã người dùng duy nhất; index role,status", "UC-ACC-01..03", "/api/auth/*; /api/users/*", "TC-AUTH-01..05; TC-USR-01..03"),
    "StudentProfile": ("Hồ sơ chính và context mục tiêu của sinh viên.", "1-1 Student; n-1 CareerGoal/Specialization.", "Không dùng status; thay goal/spec làm kết quả phân tích cũ chuyển OBSOLETE.", "studentId duy nhất; index careerGoalId,specializationId", "UC-PRO-01", "GET/PATCH /api/profiles/me", "TC-PRO-01..02"),
    "Skill": ("Định nghĩa một năng lực được dùng xuyên Assessment, Engine và Roadmap.", "n-1 SkillCategory; 1-n required skill/question/student skill.", "ACTIVE được cấu hình và đánh giá; INACTIVE chỉ đọc lịch sử.", "code và tên chuẩn hóa duy nhất; index categoryId,status", "UC-CAT-01", "/api/skills", "TC-CAT-01..02"),
    "CareerGoal": ("Định nghĩa mục tiêu nghề nghiệp và bộ yêu cầu mặc định.", "1-n Specialization/CareerRequiredSkill/StudentProfile.", "ACTIVE cho phép chọn; INACTIVE không dùng cho context mới.", "code/name duy nhất; index status", "UC-CAT-02", "/api/career-goals", "TC-CAT-03..04"),
    "Specialization": ("Chuyên sâu thuộc một mục tiêu nghề nghiệp và có thể ghi đè yêu cầu.", "n-1 CareerGoal; 1-n CareerRequiredSkill/StudentProfile.", "ACTIVE cho phép chọn trong đúng CareerGoal.", "unique careerGoalId+name; index careerGoalId,status", "UC-CAT-02", "/api/admin/specializations", "TC-CAT-03"),
    "CareerRequiredSkill": ("Ràng buộc skill, requiredLevel, trọng số, thứ tự học và giai đoạn.", "n-1 CareerGoal/Specialization/Skill.", "Không có status riêng; hiệu lực theo đối tượng cha và version cấu hình.", "unique scope+skillId; index goal/spec,phase,order", "UC-CAT-02; UC-ASM-02; UC-ENG-01", "/api/career-goals/*/required-skills", "TC-CAT-04; TC-ASM-04"),
    "AssessmentQuestion": ("Ngân hàng câu hỏi gắn skill và level.", "n-1 Skill; 1-n AssessmentOption/QuestionItem.", "DRAFT chưa dùng; ACTIVE được chọn; INACTIVE chỉ giữ lịch sử.", "index skillId,level,status,type", "UC-ASM-01", "/api/assessment-questions", "TC-ASM-01..04"),
    "AssessmentOption": ("Lựa chọn và đáp án đúng của câu hỏi.", "n-1 AssessmentQuestion; 1-n AssessmentAnswer snapshot.", "Không có status; bất biến sau khi câu hỏi đã vào assessment.", "unique questionId+orderIndex", "UC-ASM-01; UC-ASM-03", "/api/assessment-questions", "TC-ASM-01..03"),
    "Assessment": ("Một lần làm bài có hạn và policy snapshot.", "n-1 Student/CareerGoal; 1-n QuestionItem/Answer/SkillResult.", "IN_PROGRESS đang làm; GRADED đã chấm; CANCELLED hết hạn hoặc cấu hình mất hiệu lực.", "unique active attempt theo student+context; index studentId,status,expiresAt", "UC-ASM-02..04", "/api/assessments/*", "TC-ASM-04..13"),
    "AssessmentAnswer": ("Câu trả lời snapshot của sinh viên cho một item.", "n-1 Assessment/QuestionItem/Option.", "Không có status; chỉ ghi trong giao dịch submit thành công.", "unique assessmentId+questionItemId", "UC-ASM-03", "POST /api/assessments/:id/submit", "TC-ASM-08..10"),
    "AssessmentSkillResult": ("Kết quả theo skill và level sau khi chấm.", "n-1 Assessment/Skill.", "Không có status; bất biến theo lần đánh giá.", "unique assessmentId+skillId; index skillId,calculatedLevel", "UC-ASM-03; UC-ENG-01", "GET /api/assessments/me", "TC-ASM-11..13"),
    "StudentSkill": ("Mức kỹ năng tốt nhất có hiệu lực của sinh viên.", "n-1 Student/Skill; tham chiếu Evidence/AssessmentSkillResult.", "verificationStatus cho biết tự khai, chờ duyệt hay đã xác minh; không tự giảm mức.", "unique studentId+skillId; index verificationStatus,currentLevel", "UC-PRO-02; UC-EVD-02; UC-ASM-03", "/api/profiles/me/skills", "TC-PRO-03; TC-EVD-05; TC-ASM-13"),
    "Evidence": ("Yêu cầu xác minh cho một tài nguyên hồ sơ.", "n-1 Student và resource đa hình; 1-n EvidenceRevision.", "PENDING chờ duyệt; VERIFIED được chấp nhận; REJECTED phải nêu lý do.", "index studentId,status,relatedEntityType/id", "UC-EVD-01..02", "/api/evidences", "TC-EVD-01..06"),
    "EvidenceRevision": ("Lưu từng lần nộp hoặc nộp lại minh chứng.", "n-1 Evidence; tham chiếu reviewer User.", "PENDING/VERIFIED/REJECTED độc lập theo revision; một revision hiện hành.", "unique evidenceId+revisionNo; index status,createdAt", "UC-EVD-01..02", "/api/evidences/:id/revisions", "TC-EVD-06"),
    "AdvisorReview": ("Nhận xét định kỳ của cố vấn cho sinh viên được phân công.", "n-1 Student/Lecturer.", "Không có status; chỉnh sửa tạo audit hoặc revision theo policy.", "index studentId,lecturerId,createdAt", "UC-ADV-01", "/api/advisor-reviews", "TC-ADV-02"),
    "CompetencyAnalysisResult": ("Snapshot tổng hợp của một lần chạy Engine.", "n-1 Student/context/policy; 1-n CompetencyAnalysisDetail/Recommendation.", "CURRENT là kết quả hiệu lực duy nhất trong scope; OBSOLETE chỉ đọc lịch sử.", "unique current scope; index studentId,context,status,analyzedAt", "UC-ENG-01..02", "/api/competency-analysis/*", "TC-ENG-01..02"),
    "CompetencyAnalysisDetail": ("Giải thích gap, điểm đạt, trọng số và ưu tiên theo skill.", "n-1 AnalysisResult/Skill.", "Status cho biết MET/GAP/NO_DATA; bất biến theo result.", "unique analysisResultId+skillId; index status,priorityScore", "UC-ENG-01..02", "/api/competency-analysis/current/me", "ENG-AC-01..12"),
    "RecommendationTemplate": ("Mẫu hành động phát triển có điều kiện áp dụng.", "n-1 skill/category/goal/spec; n-n LearningResource.", "DRAFT chưa sinh gợi ý; ACTIVE được dùng; INACTIVE chỉ giữ lịch sử.", "unique scope+title+type; index status,priority", "UC-REC-01", "/api/recommendation-templates", "TC-REC-04"),
    "StudentRecommendation": ("Gợi ý cá nhân snapshot từ một AnalysisResult.", "n-1 Student/AnalysisResult/Skill/Template.", "NEW, SAVED, IN_PROGRESS, DONE, DISMISSED, OBSOLETE theo máy trạng thái.", "index studentId,status,priorityScore", "UC-REC-01..03", "/api/recommendations/me", "TC-REC-01..04"),
    "LearningResource": ("Tài nguyên học tập dùng cho gợi ý và lộ trình.", "n-n Skill/Template/StudentRecommendation.", "ACTIVE được chọn; INACTIVE vẫn hiển thị trong snapshot cũ.", "index type,status,title,provider", "UC-REC-02", "/api/learning-resources", "TC-REC-03"),
    "StudentResourceProgress": ("Tiến độ của sinh viên trên từng tài nguyên.", "n-1 Student/Resource/Recommendation.", "NOT_STARTED, IN_PROGRESS, COMPLETED, SKIPPED; COMPLETED là cuối.", "unique studentId+resourceId; index studentId,status", "UC-REC-02", "PATCH /api/learning-resources/:id/progress", "TC-REC-03"),
    "DevelopmentRoadmap": ("Một phiên bản lộ trình trong context nghề nghiệp.", "n-1 Student/context; 1-n RoadmapItem/Progress.", "ACTIVE, ARCHIVED, COMPLETED, CANCELLED; tối đa một ACTIVE mỗi context.", "unique studentId+version; index studentId,status", "UC-ROA-01", "/api/roadmaps/*", "TC-ROA-01..02,05"),
    "RoadmapItem": ("Nhiệm vụ có thứ tự, giai đoạn, trọng số và deadline.", "n-1 Roadmap/Skill/Recommendation; n-n dependency.", "TODO, IN_PROGRESS, DONE, CANCELLED theo điều kiện phụ thuộc.", "index roadmapId,phase,orderIndex,status", "UC-ROA-02", "PATCH /api/roadmap-items/:id/status", "TC-ROA-03..05"),
    "RoadmapItemDependency": ("Quan hệ học trước hoặc chặn giữa các mục lộ trình.", "n-1 RoadmapItem ở hai vai trò successor/predecessor.", "Không có status; hiệu lực theo roadmap version.", "unique roadmapItemId+predecessorItemId", "UC-ROA-01..02", "/api/roadmaps/*", "TC-ROA-03"),
    "StudentRoadmapProgress": ("Lịch sử tiến độ theo item và sinh viên.", "n-1 Student/Roadmap/RoadmapItem.", "Phản ánh status item tại thời điểm cập nhật; không ghi đè lịch sử quan trọng.", "index studentId,roadmapId,itemId,updatedAt", "UC-ROA-02", "PATCH /api/roadmap-items/:id/status", "TC-ROA-04..05"),
    "Notification": ("Hộp thư sự kiện theo người dùng.", "n-1 User; 1-n NotificationLog.", "isRead=false là chưa đọc; true là đã đọc, không làm mất nội dung.", "index userId,isRead,createdAt; dedup key theo sự kiện", "UC-NOT-01", "/api/notifications/me", "TC-NOT-01..02"),
    "ActivityLog": ("Nhật ký bất biến cho thao tác nhạy cảm và điều tra.", "n-1 User; tham chiếu entity đa hình.", "Không có status; chỉ append, áp dụng retention theo policy.", "index actor,action,entity,time,requestId", "UC-AUD-01", "GET /api/activity-logs", "TC-AUD-01..05"),
    "ExportJob": ("Job nền sinh PDF/XLSX và quản lý vòng đời tệp.", "n-1 owner User; 1-1 hoặc 1-n ExportLog.", "QUEUED, PROCESSING, COMPLETED, FAILED, EXPIRED.", "index ownerId,status,createdAt; idempotencyKey duy nhất", "UC-EXP-01..02", "/api/exports/*", "TC-EXP-01..03"),
    "ExportLog": ("Dấu vết yêu cầu xuất và tệp đã phát hành.", "n-1 User/ExportJob.", "Không có status riêng; giữ checksum, loại xuất và thời điểm.", "index userId,exportType,createdAt", "UC-EXP-01..02", "/api/exports/*", "TC-EXP-01..03"),
}


API_CONTRACTS = [
    ("POST", "/api/auth/login", "Xác thực và tạo phiên", "Public", "Không có scope dữ liệu", "{email,password}", "-", "200 {accessToken,user}; refresh cookie", "400/401/403/429", "Email hợp lệ; mật khẩu 8-128; User ACTIVE", "Không tiết lộ trường sai; giới hạn tần suất", "Lưu refreshTokenHash và lastLoginAt", "Không; mỗi lần tạo phiên mới", "TC-AUTH-01..03 / UC-ACC-01"),
    ("POST", "/api/auth/forgot-password", "Khởi tạo yêu cầu đặt lại mật khẩu", "Public", "Theo email chuẩn hóa", "{email}", "-", "202 accepted", "400/429", "Email đúng định dạng", "Thông điệp trung tính dù email không tồn tại", "Tạo token hash có hạn; gửi liên kết", "Idempotent trong cửa sổ chống spam", "TC-AUTH-05 / UC-ACC-02"),
    ("POST", "/api/auth/reset-password", "Đặt mật khẩu mới bằng token một lần", "Public", "Chủ token", "{token,newPassword}", "-", "200 success", "400/410/429", "Token còn hạn; mật khẩu đạt policy", "Token dùng một lần; thu hồi mọi refresh token", "Cập nhật passwordHash; ghi ActivityLog", "Token đã dùng trả cùng lỗi 410", "TC-AUTH-05 / UC-ACC-02"),
    ("PATCH", "/api/profiles/me", "Cập nhật hồ sơ và context nghề nghiệp", "Student", "Own data từ JWT", "{careerGoalId?,specializationId?,summary?,urls?}", "-", "200 StudentProfileView", "400/401/409", "Specialization thuộc goal; URL hợp lệ", "Không nhận studentId; đổi context làm analysis cũ obsolete", "Cập nhật completion; archive roadmap khi cần", "Idempotent theo payload", "TC-PRO-01..02 / UC-PRO-01"),
    ("POST", "/api/evidences", "Nộp minh chứng", "Student", "Resource của chính sinh viên", "multipart {relatedEntityType,relatedEntityId,file|url,note?}", "-", "201 EvidenceView(PENDING)", "400/403/413/415/409", "MIME,size,URL,ownership", "Mỗi lần nộp lại tạo revision; file private", "Quét file; cập nhật nguồn; gửi hàng đợi", "Idempotency-Key bắt buộc cho upload", "TC-EVD-01..02,06 / UC-EVD-01"),
    ("PATCH", "/api/evidences/:id/verify", "Xác minh minh chứng", "Advisor/Admin", "Assigned student hoặc all data", "{reviewerNote?,verifiedLevel?}", "-", "200 EvidenceView(VERIFIED)", "400/403/404/409", "Evidence PENDING; assignment ACTIVE", "verifiedLevel đúng miền; không verify lại trạng thái cuối", "Cập nhật nguồn; thông báo; ActivityLog", "Idempotent khi cùng quyết định", "TC-EVD-03,05 / UC-EVD-02"),
    ("PATCH", "/api/evidences/:id/reject", "Từ chối minh chứng", "Advisor/Admin", "Assigned student hoặc all data", "{reviewerNote}", "-", "200 EvidenceView(REJECTED)", "400/403/404/409", "reviewerNote bắt buộc", "Không từ chối ngoài assignment hoặc trạng thái cuối", "Thông báo; ActivityLog; cho phép revision mới", "Idempotent khi cùng quyết định", "TC-EVD-03..04 / UC-EVD-02"),
    ("POST", "/api/assessments/start", "Khởi tạo bài đánh giá", "Student", "Own context từ profile", "{}", "-", "201 AssessmentView+items+expiresAt", "400/401/409/429", "Coverage đủ; attempts/cooldown; không có bài còn hạn", "Snapshot policy, câu hỏi và option; tối đa 6 câu/skill", "Tạo Assessment IN_PROGRESS và QuestionItem", "Idempotency-Key bắt buộc", "TC-ASM-04..07 / UC-ASM-02"),
    ("POST", "/api/assessments/:id/submit", "Nộp và chấm bài", "Student", "Assessment của chính sinh viên", "{answers:[{questionItemId,selectedOptionId}]}", "-", "200 AssessmentResultView", "400/403/404/409", "Đủ câu; không trùng; còn hạn; option thuộc item", "Chấm nguyên tử; ngưỡng 60% liên tiếp; không hạ StudentSkill", "Ghi Answer/SkillResult; đánh dấu analysis stale", "submissionKey duy nhất", "TC-ASM-08..13 / UC-ASM-03"),
    ("GET", "/api/assessments/me", "Xem lịch sử và chi tiết đánh giá", "Student", "Own data", "-", "page,status,goalId,from,to", "200 AssessmentPage", "400/401", "Phân trang và khoảng ngày hợp lệ", "Không lộ đáp án đúng trái policy", "Không", "GET idempotent", "TC-ASM-14..16 / UC-ASM-04"),
    ("POST", "/api/competency-analysis/analyze/me", "Chạy Engine cho context hiện tại", "Student", "Own data từ JWT", "{reason?,force?}", "-", "200 CompetencyAnalysisView", "ENG-001..008; 401/409", "Context, coverage và AnalysisPolicy active", "Deterministic theo inputFingerprint+policyVersion", "Obsolete kết quả cũ; sinh gợi ý; thông báo", "Idempotency-Key + fingerprint", "TC-ENG-01..02 / UC-ENG-01"),
    ("GET", "/api/competency-analysis/current/me", "Xem kết quả phân tích hiện hành", "Student", "Own current context", "-", "include=details,recommendations", "200 CompetencyAnalysisView", "401/404", "Chỉ status CURRENT và context khớp profile", "Kết quả snapshot, không tính lại khi GET", "Không", "GET idempotent", "ENG-AC-09..12 / UC-ENG-02"),
    ("GET", "/api/recommendations/me", "Liệt kê gợi ý cá nhân", "Student", "Own current context", "-", "status,skillId,page", "200 RecommendationPage", "400/401", "Filter hợp lệ", "Mặc định ẩn OBSOLETE trừ khi yêu cầu lịch sử", "Không", "GET idempotent", "TC-REC-01..04 / UC-REC-01"),
    ("PATCH", "/api/recommendations/:id/status", "Chuyển trạng thái gợi ý", "Student", "Own recommendation", "{status,note?}", "-", "200 RecommendationView", "400/403/404/409", "State transition hợp lệ", "DONE/DISMISSED/OBSOLETE là cuối theo policy", "Đồng bộ resource/roadmap; ActivityLog khi mở lại", "transitionKey duy nhất", "TC-REC-01..02 / UC-REC-01"),
    ("PATCH", "/api/learning-resources/:id/progress", "Cập nhật tiến độ tài nguyên", "Student", "Own progress", "{status,note?}", "-", "200 ResourceProgressView", "400/403/404/409", "State transition hợp lệ", "COMPLETED là cuối; SKIPPED có thể quay lại", "Đồng bộ recommendation và roadmap item", "transitionKey duy nhất", "TC-REC-03 / UC-REC-02"),
    ("POST", "/api/roadmaps/generate/me", "Sinh lộ trình cá nhân", "Student", "Own current context", "{archiveCurrent?:true}", "-", "201 RoadmapView", "400/401/409", "Có analysis CURRENT và recommendation actionable", "Một ACTIVE/context; áp dụng dependency và phase", "Archive bản cũ; tạo item/progress", "Idempotency-Key bắt buộc", "TC-ROA-01..02 / UC-ROA-01"),
    ("GET", "/api/roadmaps/current/me", "Xem lộ trình hiện hành", "Student", "Own data", "-", "include=items,dependencies,progress", "200 RoadmapView", "401/404", "Chỉ ACTIVE của context hiện tại", "Không trộn version hoặc context", "Không", "GET idempotent", "TC-ROA-01 / UC-ROA-01"),
    ("PATCH", "/api/roadmap-items/:id/status", "Cập nhật mục lộ trình", "Student", "Item thuộc roadmap của mình", "{status,note?}", "-", "200 RoadmapItemView", "400/403/404/409", "Dependency và transition hợp lệ", "Mở lại DONE cần note; CANCELLED là cuối", "Ghi progress; tính roadmap; đồng bộ gợi ý", "transitionKey duy nhất", "TC-ROA-03..05 / UC-ROA-02"),
    ("GET", "/api/dashboard/student", "Dashboard sinh viên", "Student", "Own current context", "-", "-", "200 StudentDashboardView", "401", "Context nhất quán", "Chỉ số từ snapshot CURRENT", "Không", "GET idempotent", "TC-DSH-01 / UC-DSH-01"),
    ("GET", "/api/dashboard/lecturer", "Dashboard cố vấn", "Advisor/Lecturer", "Assigned students ACTIVE", "-", "filters,page", "200 LecturerDashboardView", "401/403", "Assignment được lọc trong truy vấn", "Không lọc sau khi tải toàn bộ", "Không", "GET idempotent", "TC-DSH-02 / UC-DSH-01"),
    ("GET", "/api/dashboard/admin", "Dashboard quản trị", "Admin", "All data", "-", "from,to,goalId", "200 AdminDashboardView", "400/401/403", "Khoảng ngày hợp lệ", "Không trả PII không cần thiết", "Không", "GET idempotent", "TC-DSH-03 / UC-DSH-01"),
    ("GET", "/api/notifications/me", "Danh sách thông báo", "Authenticated", "Own notifications", "-", "isRead,type,page", "200 NotificationPage", "400/401", "Pagination hợp lệ", "Chỉ userId từ JWT", "Không", "GET idempotent", "TC-NOT-01..02 / UC-NOT-01"),
    ("PATCH", "/api/notifications/:id/read", "Đánh dấu đã đọc", "Authenticated", "Own notification", "{}", "-", "200 NotificationView", "401/404", "Notification thuộc user", "Không tiết lộ tồn tại ngoài scope", "Cập nhật readAt", "Idempotent", "TC-NOT-02 / UC-NOT-01"),
    ("POST", "/api/exports/portfolio", "Xuất portfolio PDF", "Student", "Own profile", "{sections,locale,privacy}", "-", "202 ExportJobView", "400/401/413/429", "Section và privacy hợp lệ", "Snapshot dữ liệu; URL ký có hạn", "Queue job; ExportLog; notification", "Idempotency-Key bắt buộc", "TC-EXP-01 / UC-EXP-01"),
    ("POST", "/api/exports/report", "Xuất báo cáo XLSX", "Advisor/Admin", "Assigned students hoặc all data", "{type,filters,columns}", "-", "202 ExportJobView", "400/403/413/429", "Filter/column theo loại báo cáo", "Scope áp dụng trong query; giới hạn dòng", "Queue job; ExportLog; notification", "Idempotency-Key bắt buộc", "TC-EXP-02..03 / UC-EXP-02"),
]


SCREEN_ROWS = [
    ["Đăng nhập", "Mọi vai trò", "Tạo phiên và điều hướng", "email,password", "Đăng nhập", "POST /api/auth/login", "spinner; lỗi 401/403/429"],
    ["Dashboard sinh viên", "Student", "Tóm tắt hồ sơ, gap và việc tiếp theo", "không", "Mở profile/assessment/analysis/roadmap", "GET /api/dashboard/student", "skeleton; chưa chọn goal; lỗi context"],
    ["Thiết lập hồ sơ", "Student", "Cập nhật profile và mục tiêu", "goal,spec,summary,URL", "Lưu; tải avatar", "PATCH /api/profiles/me", "loading; validation; conflict"],
    ["Minh chứng", "Student", "Nộp và theo dõi xác minh", "resource,file|URL,note", "Nộp; nộp lại; xem", "POST /api/evidences", "upload progress; empty; rejected"],
    ["Assessment", "Student", "Làm bài có thời hạn", "answers", "Bắt đầu; chọn; nộp", "/api/assessments/*", "coverage; cooldown; timeout; submit conflict"],
    ["Kết quả phân tích", "Student", "Giải thích match, risk và gap", "reason?", "Phân tích lại; mở gợi ý", "/api/competency-analysis/*", "processing; no-data; ENG error"],
    ["Gợi ý", "Student", "Chọn hành động phát triển", "status,note", "Lưu; bắt đầu; hoàn thành; bỏ qua", "/api/recommendations/me", "empty; obsolete; transition error"],
    ["Lộ trình", "Student", "Thực hiện nhiệm vụ theo giai đoạn", "status,note", "Sinh; bắt đầu; hoàn thành; mở lại", "/api/roadmaps/*", "no-analysis; dependency; overdue"],
    ["Dashboard cố vấn", "Advisor", "Ưu tiên sinh viên cần hỗ trợ", "filters", "Mở sinh viên/hàng đợi", "GET /api/dashboard/lecturer", "empty assignment; permission"],
    ["Chi tiết sinh viên", "Advisor", "Xem profile, analysis và roadmap", "studentId từ route", "Mở minh chứng; tạo nhận xét", "/api/profiles/:studentId", "404 an toàn; assignment inactive"],
    ["Duyệt minh chứng", "Advisor/Admin", "Verify hoặc reject đúng scope", "reviewerNote,verifiedLevel?", "Xác minh; từ chối", "/api/evidences/:id/*", "preview error; conflict; forbidden"],
    ["Nhận xét cố vấn", "Advisor", "Ghi đánh giá định kỳ", "reviewContent,suggestion,rating", "Gửi nhận xét", "/api/advisor-reviews", "validation; forbidden"],
    ["Dashboard quản trị", "Admin", "Giám sát hệ thống và chất lượng cấu hình", "filters", "Đi tới catalog/policy/report", "GET /api/dashboard/admin", "degraded metrics; error"],
    ["Danh mục kỹ năng", "Admin", "Quản lý category và skill", "code,name,level,status", "CRUD; activate; deactivate", "/api/skill-categories; /api/skills", "duplicate; in-use; validation"],
    ["Mục tiêu nghề nghiệp", "Admin", "Quản lý goal/spec/required skill", "goal,spec,skill,level,weight", "CRUD; mô phỏng; kích hoạt", "/api/career-goals/*", "coverage warning; conflict"],
    ["Ngân hàng câu hỏi", "Admin", "Bảo đảm coverage", "skill,level,type,options", "CRUD; filter gap; activate", "/api/assessment-questions", "invalid options; missing coverage"],
    ["Policy", "Admin", "Phiên bản hóa Assessment/Engine policy", "thresholds,limits,weights", "Draft; simulate; activate", "/api/analysis-policies", "validation; concurrent activation"],
    ["Thông báo", "Authenticated", "Đọc sự kiện cần hành động", "filters", "Mở; đánh dấu; đọc tất cả", "/api/notifications/me", "empty; pagination; 404"],
    ["Portfolio/Báo cáo", "Student/Advisor/Admin", "Tạo và tải file có hạn", "sections/filters/columns", "Tạo job; theo dõi; tải", "/api/exports/*", "queued; failed; expired"],
]


STATE_MACHINES = {
    "Assessment": [
        ["IN_PROGRESS", "Bài đang làm và còn hạn", "GRADED, CANCELLED", "Student/Hệ thống", "Submit hợp lệ hoặc hết hạn/cấu hình mất hiệu lực", "Chấm nguyên tử hoặc giải phóng bài"],
        ["GRADED", "Đã chấm và lưu kết quả", "-", "Hệ thống", "Giao dịch submit hoàn tất", "Cập nhật StudentSkill; đánh dấu analysis stale"],
        ["CANCELLED", "Bài không còn hợp lệ", "-", "Hệ thống/Admin", "Quá hạn hoặc question/policy bị vô hiệu", "Không chấm; ghi lý do"],
    ],
    "Evidence": [
        ["PENDING", "Đang chờ xác minh", "VERIFIED, REJECTED", "Advisor/Admin", "Assignment ACTIVE và resource hợp lệ", "Thông báo kết quả; cập nhật nguồn"],
        ["VERIFIED", "Minh chứng được chấp nhận", "-", "Advisor/Admin", "Có quyết định xác minh", "Ghi verifier, verifiedAt, ActivityLog"],
        ["REJECTED", "Minh chứng bị từ chối", "PENDING qua revision mới", "Student", "Có reviewerNote; sinh viên nộp lại", "Tạo EvidenceRevision, giữ lịch sử"],
    ],
    "AnalysisResult": [
        ["CURRENT", "Kết quả hiệu lực trong context", "OBSOLETE", "Hệ thống", "Context/policy/input thay đổi hoặc chạy lại", "Ẩn khỏi mặc định; recommendation cũ obsolete"],
        ["OBSOLETE", "Snapshot lịch sử chỉ đọc", "-", "Hệ thống", "Có result CURRENT thay thế", "Không dùng để sinh roadmap mới"],
    ],
    "Recommendation": [
        ["NEW", "Gợi ý mới chưa được chọn", "SAVED, DISMISSED", "Student", "Context còn hiệu lực", "Ghi thời điểm lựa chọn"],
        ["SAVED", "Đã lưu để thực hiện", "IN_PROGRESS, DISMISSED", "Student", "Có resource hoặc hành động hợp lệ", "Có thể sinh roadmap item"],
        ["IN_PROGRESS", "Đang thực hiện", "DONE, SAVED", "Student/Hệ thống", "Tiến độ resource thay đổi", "Đồng bộ roadmap"],
        ["DONE", "Đã hoàn thành", "-", "Student/Hệ thống", "Đủ điều kiện hoàn thành", "Đóng item liên quan"],
        ["DISMISSED", "Bị bỏ qua", "-", "Student", "Có lý do tùy policy", "Loại khỏi lộ trình mới"],
        ["OBSOLETE", "Không còn phù hợp context", "-", "Hệ thống", "Analysis hoặc catalog thay đổi", "Chỉ đọc lịch sử"],
    ],
    "LearningResourceProgress": [
        ["NOT_STARTED", "Chưa bắt đầu", "IN_PROGRESS, SKIPPED", "Student", "Resource còn truy cập được", "Ghi startedAt khi bắt đầu"],
        ["IN_PROGRESS", "Đang học", "COMPLETED, SKIPPED", "Student", "Transition hợp lệ", "Cập nhật lastAccessedAt"],
        ["COMPLETED", "Đã hoàn thành", "-", "Student/Hệ thống", "Student xác nhận hoặc provider callback", "Đồng bộ recommendation/roadmap"],
        ["SKIPPED", "Tạm bỏ qua", "IN_PROGRESS", "Student", "Có thể quay lại", "Không tính là hoàn thành"],
    ],
    "Roadmap": [
        ["ACTIVE", "Lộ trình hiện hành", "ARCHIVED, COMPLETED, CANCELLED", "Student/Hệ thống", "Sinh lại, hoàn thành mọi item hoặc hủy", "Chốt progress và version"],
        ["ARCHIVED", "Phiên bản lịch sử", "-", "Hệ thống", "Có ACTIVE mới", "Chỉ đọc"],
        ["COMPLETED", "Hoàn thành mục tiêu lộ trình", "-", "Hệ thống", "Mọi item không hủy đã DONE", "progress=100; thông báo"],
        ["CANCELLED", "Lộ trình bị hủy", "-", "Student/Admin", "Có lý do hợp lệ", "Dừng item chưa hoàn tất"],
    ],
    "RoadmapItem": [
        ["TODO", "Chưa bắt đầu", "IN_PROGRESS, CANCELLED", "Student", "Dependency đã DONE", "Ghi startedAt"],
        ["IN_PROGRESS", "Đang thực hiện", "TODO, DONE, CANCELLED", "Student", "Transition hợp lệ", "Tính lại roadmap progress"],
        ["DONE", "Đã hoàn thành", "IN_PROGRESS", "Student", "Mở lại bắt buộc note", "Đồng bộ recommendation"],
        ["CANCELLED", "Không còn thực hiện", "-", "Student/Admin", "Có lý do", "Loại khỏi mẫu số tiến độ theo policy"],
    ],
    "Notification": [
        ["UNREAD", "Chưa được người dùng mở", "READ", "Owner", "Mark one/all hoặc mở chi tiết", "Ghi readAt"],
        ["READ", "Đã đọc", "-", "Owner", "Đã có readAt", "Không xóa nội dung"],
    ],
    "ExportJob": [
        ["QUEUED", "Đang chờ worker", "PROCESSING, FAILED", "Hệ thống", "Job hợp lệ và chưa hết hạn", "Ghi queuedAt"],
        ["PROCESSING", "Đang sinh file", "COMPLETED, FAILED", "Worker", "Lock job thành công", "Ghi progress và heartbeat"],
        ["COMPLETED", "File sẵn sàng", "EXPIRED", "Worker/Scheduler", "Checksum và URL đã tạo", "Thông báo owner"],
        ["FAILED", "Sinh file thất bại", "QUEUED", "Worker/Admin", "Retry còn hạn", "Ghi errorCode, retryCount"],
        ["EXPIRED", "File đã bị thu hồi", "-", "Scheduler", "expiresAt đã qua", "Xóa vật lý; giữ ExportLog"],
    ],
}


def element_style(el, doc: Document) -> str:
    if el.tag != qn("w:p"):
        return ""
    p = Paragraph(el, doc)
    return p.style.name if p.style else ""


def append_element(doc: Document, el) -> None:
    body = doc.element.body
    sect_pr = body.sectPr
    body.insert(list(body).index(sect_pr), deepcopy(el))


def append_range(doc: Document, source_elements, start: int, end: int, skip_heading1=True) -> None:
    for el in source_elements[start:end]:
        if skip_heading1 and element_style(el, doc) == "Heading 1":
            continue
        append_element(doc, el)


def format_table(table, widths, font_size=8.0, header_fill=LIGHT_GRAY) -> None:
    table.style = "Table Grid"
    table.alignment = 0
    set_table_geometry(table, widths)
    set_table_borders(table)
    for ridx, row in enumerate(table.rows):
        set_cant_split(row)
        if ridx == 0:
            set_repeat_header(row)
        for cidx, cell in enumerate(row.cells):
            set_cell_margins(cell, 75, 100, 75, 100)
            cell.vertical_alignment = 1
            if ridx == 0:
                shade_cell(cell, header_fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if ridx == 0 or (len(widths) <= 3 and cidx == 0):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=(ridx == 0))


def add_plain_table(doc: Document, headers, rows, widths, font_size=8.0) -> Table:
    table = doc.add_table(rows=1, cols=len(headers))
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = str(value)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = str(value)
    format_table(table, widths, font_size)
    return table


def add_toc(doc: Document) -> None:
    table = add_plain_table(
        doc,
        ["Nội dung", "Trang"],
        [[f"{i}. {title}", "—"] for i, title in enumerate(CHAPTERS, 1)],
        [8200, 1160],
        9.0,
    )
    for row in table.rows[1:]:
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_page_break()


def function_level3_rows():
    rows = []
    for level1, group, level2, subgroup, details in FUNCTION_TREE:
        actor, input_text, output_text, rules = GROUP_META[level1]
        for item in [x.strip() for x in details.split(";") if x.strip()]:
            match = re.match(r"([A-Z]\d+(?:\.\d+){2})\s+(.+)", item)
            if not match:
                continue
            code, name = match.groups()
            desc = f"Thực hiện thao tác {name.lower()} trong nhóm {subgroup.lower()}, bảo đảm đúng role và scope dữ liệu."
            rows.append([code, name, actor, desc, input_text, output_text, rules])
    return rows


def add_function_design(w: Writer) -> None:
    w.heading("Đặc tả chức năng cấp 3", 2)
    w.paragraph(
        "Mỗi chức năng cấp 3 là đơn vị nhỏ nhất được dùng để lập use case, API, màn hình và test. "
        "Actor trong bảng là actor được phép khởi tạo; các side effect tự động do hệ thống thực hiện."
    )
    w.table(
        ["Mã", "Tên chức năng", "Actor", "Mô tả", "Đầu vào", "Đầu ra", "Quy tắc"],
        function_level3_rows(),
        [800, 1400, 1000, 2100, 1200, 1200, 1660],
        "Đặc tả chức năng cấp 3",
        6.6,
    )


def uc_rule(uc_id: str) -> str:
    prefix = uc_id.split("-")[1]
    return {
        "ACC": "BR-ACC-01..08",
        "CAT": "BR-CAT-01..10",
        "PRO": "BR-PRO-01..08",
        "EVD": "BR-EVD-01..10",
        "ASM": "BR-ASM-01..16",
        "ENG": "BR-ENG-01..12",
        "REC": "BR-REC-01..12",
        "ROA": "BR-ROA-01..12",
        "ADV": "BR-ADV-01..06",
        "DSH": "BR-DSH-01..08",
        "NOT": "BR-NOT-01..08",
        "EXP": "BR-EXP-01..10",
        "AUD": "BR-AUD-01..08",
    }.get(prefix, "BR-GEN-01")


def clean_design_phrase(text: str) -> str:
    replacements = {
        "API dự kiến/hiện có": "API liên quan",
        " (đề xuất)": "",
        " dự kiến": "",
        "hiện có/": "",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def enhance_use_case_tables(doc: Document) -> None:
    targets = []
    current_uc = None
    current_name = None
    for el in list(doc.element.body):
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            if p.style and p.style.name == "Heading 3":
                m = re.search(r"(UC-[A-Z]+-\d+)\s*-\s*(.+)$", p.text.strip())
                if m:
                    current_uc, current_name = m.group(1), m.group(2)
        elif el.tag == qn("w:tbl") and current_uc:
            table = Table(el, doc)
            if len(table.columns) == 2 and table.cell(0, 0).text.strip() == "Thuộc tính":
                targets.append((el, current_uc, current_name))
                current_uc = current_name = None

    for old_el, uc_id, name in targets:
        old = Table(old_el, doc)
        data = {row.cells[0].text.strip(): clean_design_phrase(row.cells[1].text.strip()) for row in old.rows[1:]}
        if "Mã use case" in data:
            continue
        actor = data.get("Actor", "Người dùng/Hệ thống")
        pre = data.get("Tiền điều kiện", "Đã xác thực và có quyền phù hợp.")
        post = data.get("Hậu điều kiện", "Dữ liệu được cập nhật nhất quán.")
        main = data.get("Luồng chính", "Actor gửi yêu cầu; hệ thống kiểm tra quyền, validation, quy tắc; lưu và trả kết quả.")
        exception = data.get("Ngoại lệ và quy tắc", data.get("Luồng ngoại lệ", "Từ chối khi sai quyền, validation hoặc trạng thái."))
        related_data = data.get("Dữ liệu", "Các thực thể nghiệp vụ liên quan")
        api = data.get("API liên quan", data.get("API dự kiến/hiện có", "Theo hợp đồng API Chương 11"))
        tests = data.get("Kiểm thử", "Theo bộ test Chương 14")
        goal = f"Cho phép {actor.lower()} thực hiện {name.lower()} với kết quả có thể kiểm tra và truy vết."
        alternative = (
            "Nếu bản ghi hoặc kết quả hợp lệ đã tồn tại, hệ thống trả dữ liệu hiện hành hoặc yêu cầu actor xác nhận thay thế; "
            "actor có thể hủy trước bước xác nhận cuối."
        )
        rows = [
            ["Mã use case", uc_id], ["Tên use case", name], ["Actor", actor], ["Mục tiêu", goal],
            ["Tiền điều kiện", pre], ["Hậu điều kiện", post], ["Luồng chính", main],
            ["Luồng thay thế", alternative], ["Luồng ngoại lệ", exception],
            ["Dữ liệu vào", f"Payload/query và định danh liên quan đến {related_data}."],
            ["Dữ liệu ra", post], ["Business rule liên quan", uc_rule(uc_id)],
            ["API liên quan", api], ["Test case liên quan", tests],
        ]
        new_table = add_plain_table(doc, ["Thuộc tính", "Đặc tả"], rows, [1900, 7460], 8.4)
        old_el.addprevious(new_table._tbl)
        old_el.getparent().remove(old_el)


def add_full_uc(w: Writer, uc_id, name, actor, goal, pre, post, main, alt, exc, data_in, data_out, rules, api, tests):
    w.heading(f"{uc_id} - {name}", 3)
    w.table(
        ["Thuộc tính", "Đặc tả"],
        [["Mã use case", uc_id], ["Tên use case", name], ["Actor", actor], ["Mục tiêu", goal],
         ["Tiền điều kiện", pre], ["Hậu điều kiện", post], ["Luồng chính", main],
         ["Luồng thay thế", alt], ["Luồng ngoại lệ", exc], ["Dữ liệu vào", data_in],
         ["Dữ liệu ra", data_out], ["Business rule liên quan", rules], ["API liên quan", api],
         ["Test case liên quan", tests]],
        [1900, 7460],
        f"Đặc tả {uc_id}",
        8.4,
    )


def add_missing_use_cases(w: Writer) -> None:
    w.heading("Use case bổ sung cho luồng phân tích và thông báo", 2)
    add_full_uc(
        w, "UC-ENG-02", "Xem kết quả phân tích hiện hành", "Sinh viên/Cố vấn",
        "Trình bày match, coverage, risk, gap và giải thích theo đúng context.",
        "Có AnalysisResult CURRENT; cố vấn có assignment ACTIVE.",
        "Kết quả snapshot được hiển thị mà không chạy lại Engine.",
        "Xác định scope; tải result và details; kiểm tra context; trả biểu đồ, giải thích và liên kết gợi ý.",
        "Nếu không có CURRENT, hiển thị hành động hoàn thiện hồ sơ hoặc chạy phân tích.",
        "Ngoài scope trả 404 an toàn; dữ liệu snapshot thiếu trả ENG-007.",
        "studentId từ phiên hoặc assignment; include options.", "CompetencyAnalysisView.",
        "BR-ENG-09..12", "GET /api/competency-analysis/current/me", "ENG-AC-09..12",
    )
    add_full_uc(
        w, "UC-REC-03", "Sinh gợi ý phát triển", "Hệ thống",
        "Chuyển các gap có ưu tiên thành hành động và tài nguyên phù hợp.",
        "Có AnalysisResult CURRENT và RecommendationTemplate ACTIVE.",
        "StudentRecommendation NEW được tạo theo snapshot.",
        "Lọc detail GAP; xếp priority; match template theo skill/context/level; snapshot resource; lưu kết quả.",
        "Nếu không có template theo skill, dùng template category rồi template chung.",
        "Không tạo gợi ý trùng fingerprint; không dùng template INACTIVE.",
        "AnalysisResult, detail, template, resource.", "Danh sách StudentRecommendation NEW.",
        "BR-REC-01..06", "Side effect của POST /api/competency-analysis/analyze/me", "TC-REC-04",
    )
    add_full_uc(
        w, "UC-NOT-02", "Phát thông báo từ sự kiện miền", "Hệ thống",
        "Tạo thông báo đúng người nhận, không trùng và có liên kết hành động.",
        "Sự kiện miền hợp lệ; người nhận tồn tại.", "Notification UNREAD và NotificationLog được ghi.",
        "Nhận event; xác định recipient; áp dụng preference/dedup; render nội dung; lưu; gửi theo channel.",
        "Nếu channel ngoài ứng dụng thất bại, giữ IN_APP và đưa log vào retry.",
        "Event thiếu entity hoặc recipient bị bỏ và ghi lỗi vận hành; không tạo bản sao trong dedup window.",
        "eventType, recipient, entity, template data.", "Notification và delivery log.",
        "BR-NOT-01..08", "Domain event nội bộ; GET/PATCH /api/notifications/* để đọc", "TC-NOT-01..02",
    )


def add_skill_scale(w: Writer) -> None:
    w.heading("Thang kỹ năng dùng chung 0-5", 2)
    levels = [
        ["0", "Chưa có dữ liệu/chưa đạt", "Chưa có bằng chứng đáng tin cậy hoặc không vượt câu hỏi level 1; cần học từ nền tảng."],
        ["1", "Nhận biết", "Nhận diện thuật ngữ, mục đích và ví dụ; thực hiện được thao tác khi có hướng dẫn từng bước."],
        ["2", "Cơ bản", "Hoàn thành tác vụ chuẩn, phạm vi nhỏ; giải thích được lựa chọn và xử lý lỗi phổ biến với tài liệu hỗ trợ."],
        ["3", "Thực hành", "Tự triển khai tác vụ thực tế đầu-cuối, kiểm thử kết quả và xử lý ngoại lệ thông thường."],
        ["4", "Thành thạo", "Thiết kế giải pháp, tối ưu chất lượng, review người khác và xử lý tình huống phức tạp."],
        ["5", "Nâng cao", "Định hình chuẩn, kiến trúc hoặc chiến lược; giải quyết bài toán mới và dẫn dắt cải tiến ở quy mô lớn."],
    ]
    w.table(["Level", "Tên mức", "Hành vi quan sát được"], levels, [900, 2100, 6360], "Thang kỹ năng 0-5 dùng chung", 8.2)
    samples = {
        "React": ["Chưa nhận diện component", "Đọc JSX và props", "Tạo component/form cơ bản", "Quản lý state, routing và gọi API", "Tối ưu render, test và kiến trúc component", "Thiết kế nền tảng frontend và chuẩn dùng chung"],
        "SQL": ["Chưa viết truy vấn", "Hiểu bảng, cột, SELECT", "JOIN, GROUP BY và CRUD chuẩn", "Thiết kế truy vấn nhiều bảng, transaction", "Tối ưu plan, index và tính nhất quán", "Thiết kế chiến lược dữ liệu ở quy mô lớn"],
        "Git": ["Chưa dùng repository", "Clone, status, commit", "Branch, merge, pull request cơ bản", "Rebase, resolve conflict và quy trình nhóm", "Thiết kế branching/release và review lịch sử", "Chuẩn hóa DevEx, governance và phục hồi sự cố"],
        "Problem Solving": ["Chưa hình thành cách tiếp cận", "Nêu lại bài toán và dữ kiện", "Chia nhỏ và chọn giải pháp quen thuộc", "So sánh phương án, kiểm chứng và xử lý edge case", "Mô hình hóa bài toán phức tạp và hướng dẫn người khác", "Tạo phương pháp mới và dẫn dắt quyết định liên ngành"],
        "NestJS/Backend API": ["Chưa hiểu request-response", "Nhận biết controller/service/DTO", "Tạo CRUD có validation", "Thiết kế module, auth, transaction và test", "Thiết kế contract, hiệu năng, observability và bảo mật", "Dẫn dắt kiến trúc dịch vụ và chuẩn API toàn hệ thống"],
    }
    rows = [[skill, str(level), behavior] for skill, behaviors in samples.items() for level, behavior in enumerate(behaviors)]
    w.table(["Kỹ năng mẫu", "Level", "Hành vi quan sát được"], rows, [2100, 800, 6460], "Thang mức theo kỹ năng mẫu", 7.5)


def add_question_coverage(w: Writer) -> None:
    w.heading("Question Coverage Policy", 2)
    w.callout(
        "Điều kiện chặn tạo bài",
        "Với mỗi required skill, ngân hàng phải có câu hỏi ACTIVE ở mọi level từ 1 đến requiredLevel. "
        "Thiếu bất kỳ cặp skill-level nào thì POST /api/assessments/start trả lỗi coverage, nêu rõ skillId và missingLevels; không tạo Assessment dở dang.",
        LIGHT_BLUE,
    )
    policy = [
        ["Coverage tối thiểu", "2 câu ACTIVE cho mỗi skill-level bắt buộc", "Cho phép chọn ngẫu nhiên và giảm lặp giữa các attempt"],
        ["Số câu chọn", "Tối thiểu 1 câu/level; tối đa 6 câu/skill", "Tổng câu phụ thuộc số skill nhưng không vượt 60"],
        ["Thời gian", "60 phút từ startedAt", "expiresAt được snapshot; quá hạn chuyển CANCELLED và không chấm"],
        ["Cooldown", "30 phút sau attempt kết thúc", "Response trả retryAt"],
        ["Max attempts", "3 lần cho student + context + policyVersion", "Đổi policyVersion chỉ mở lại nếu quy định cho phép"],
        ["Ngưỡng qua level", ">=60% số câu của level", "Tỷ lệ dùng số câu snapshot, không dùng ngân hàng hiện tại"],
        ["Đạt liên tiếp", "Level k chỉ đạt khi mọi level 1..k đều qua", "Level cao đạt nhưng level thấp trượt thì calculatedLevel=level trước chuỗi, có thể là 0"],
        ["Chống submit trùng", "submissionKey duy nhất và khóa Assessment khi chấm", "Lần lặp cùng key trả kết quả cũ; key khác sau GRADED trả 409"],
        ["Snapshot", "Lưu questionId, text, type, level, option text/order/isCorrect và policyVersion", "Chỉnh ngân hàng sau khi bắt đầu không làm đổi bài"],
        ["Giao dịch", "Answer, SkillResult, status và StudentSkill trong một transaction", "Lỗi bất kỳ rollback toàn bộ"],
    ]
    w.table(["Thành phần", "Chính sách", "Quy tắc triển khai"], policy, [1900, 3100, 4360], "Question Coverage Policy", 7.8)


def add_entity_fields(w: Writer, name: str, fields) -> None:
    w.heading(f"Bảng {name}", 3)
    w.table(
        ["Field", "Kiểu", "Bắt buộc", "Ràng buộc/unique/index", "Ý nghĩa nghiệp vụ"],
        fields,
        [1400, 1150, 900, 2550, 3360],
        f"Từ điển dữ liệu {name}",
        7.2,
    )


def add_missing_entities(w: Writer) -> None:
    w.heading("Các bảng bắt buộc của thiết kế đích", 2)
    add_entity_fields(w, "EvidenceRevision", [
        ["id", "UUID", "Có", "PK", "Định danh revision."], ["evidenceId", "UUID", "Có", "FK Evidence; index", "Minh chứng gốc."],
        ["revisionNo", "Int", "Có", "unique evidenceId+revisionNo", "Số lần nộp tăng đơn điệu."], ["fileUrl", "String?", "Không", "private URL", "Tệp của revision."],
        ["evidenceUrl", "String?", "Không", "URL hợp lệ", "Liên kết thay cho tệp."], ["checksum", "String?", "Không", "index", "Phát hiện trùng và kiểm tra toàn vẹn."],
        ["status", "EvidenceStatus", "Có", "index", "PENDING/VERIFIED/REJECTED của revision."], ["reviewerId", "UUID?", "Không", "FK User", "Người ra quyết định."],
        ["reviewerNote", "Text?", "Không", "Bắt buộc khi REJECTED", "Lý do hoặc ghi chú xác minh."], ["createdAt", "DateTime", "Có", "index", "Thời điểm nộp revision."],
    ])
    add_entity_fields(w, "RoadmapItemDependency", [
        ["roadmapItemId", "UUID", "Có", "FK RoadmapItem", "Mục bị phụ thuộc."], ["predecessorItemId", "UUID", "Có", "FK RoadmapItem", "Mục phải hoàn thành trước."],
        ["dependencyType", "Enum", "Có", "HARD/SOFT", "HARD chặn bắt đầu; SOFT chỉ cảnh báo."], ["createdAt", "DateTime", "Có", "-", "Thời điểm tạo quan hệ."],
        ["(constraint)", "-", "Có", "unique item+predecessor; check khác nhau", "Không tạo vòng lặp hoặc tự phụ thuộc."],
    ])
    add_entity_fields(w, "ExportJob", [
        ["id", "UUID", "Có", "PK", "Định danh job."], ["ownerId", "UUID", "Có", "FK User; index", "Chủ sở hữu kết quả xuất."],
        ["exportType", "Enum", "Có", "index", "PORTFOLIO_PDF hoặc REPORT_XLSX."], ["filters", "JSON", "Có", "snapshot", "Bộ lọc đã áp dụng scope."],
        ["status", "ExportJobStatus", "Có", "index", "QUEUED/PROCESSING/COMPLETED/FAILED/EXPIRED."], ["progress", "Int", "Có", "0..100", "Phần trăm xử lý."],
        ["fileUrl", "String?", "Không", "signed/private", "Đường dẫn tải có hạn."], ["checksum", "String?", "Không", "-", "Kiểm tra toàn vẹn tệp."],
        ["idempotencyKey", "String", "Có", "unique ownerId+key", "Chống tạo job trùng."], ["errorCode", "String?", "Không", "index", "Mã lỗi ổn định cho client."],
        ["expiresAt", "DateTime?", "Không", "index", "Thời điểm thu hồi file."], ["createdAt", "DateTime", "Có", "index", "Thời điểm tạo job."],
    ])


def find_entity_targets(doc: Document):
    targets = []
    current_name = None
    current_heading = None
    for el in list(doc.element.body):
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            if p.style and p.style.name == "Heading 3":
                m = re.search(r"Bảng\s+([A-Za-z][A-Za-z0-9]+)$", p.text.strip())
                if m:
                    current_name = m.group(1)
                    current_heading = el
        elif el.tag == qn("w:tbl") and current_name:
            table = Table(el, doc)
            headers = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
            if headers and headers[0] in {"Trường", "Field"}:
                targets.append((current_name, current_heading, el))
                current_name = current_heading = None
    return targets


def insert_entity_contexts(doc: Document) -> None:
    for name, heading_el, field_table_el in find_entity_targets(doc):
        if name not in ENTITY_META:
            continue
        purpose, relations, states, unique_index, use_cases, apis, tests = ENTITY_META[name]
        caption = doc.add_paragraph(f"Bảng 0. Ngữ cảnh nghiệp vụ {name}", style="Caption")
        caption.paragraph_format.keep_with_next = True
        table = add_plain_table(
            doc,
            ["Thuộc tính thiết kế", "Đặc tả"],
            [["Mục đích nghiệp vụ", purpose], ["Quan hệ chính", relations], ["Trạng thái và ý nghĩa", states],
             ["Unique/index", unique_index], ["Use case tạo/sửa/xóa", use_cases], ["API liên quan", apis],
             ["Test case liên quan", tests]],
            [2050, 7310],
            7.8,
        )
        heading_el.addnext(table._tbl)
        heading_el.addnext(caption._p)

        field_table = Table(field_table_el, doc)
        state_text = states
        if len(field_table.columns) >= 5:
            for row in field_table.rows[1:]:
                field_name = row.cells[0].text.strip()
                if field_name in {"status", "verificationStatus", "isRead"}:
                    row.cells[-1].text = state_text
                    for p in row.cells[-1].paragraphs:
                        for run in p.runs:
                            set_run_font(run, size=7.2)


def add_api_design(w: Writer) -> None:
    w.heading("Quy ước hợp đồng", 2)
    w.bullets([
        "Base path /api; JSON UTF-8; multipart chỉ cho upload; thời gian ISO 8601 UTC; mã lỗi nghiệp vụ ổn định.",
        "Response thành công dùng {success:true,data,meta?}; response lỗi dùng {success:false,statusCode,code,message,path,timestamp,requestId}.",
        "Role guard chỉ là lớp đầu; service phải áp dụng own/assigned/all scope ngay trong truy vấn database.",
        "Idempotency-Key bắt buộc với upload, start assessment, submit, chạy Engine, sinh roadmap và tạo export job.",
        "Mọi mutation nhạy cảm ghi ActivityLog với actor, subject, reason, entity, requestId và kết quả.",
    ])
    catalog = [[m, e, role, scope, purpose, tests.split("/")[0].strip()] for m, e, purpose, role, scope, *_rest, tests in API_CONTRACTS]
    w.table(["Method", "Endpoint", "Role", "Scope", "Mục đích", "Test"], catalog,
            [700, 2700, 1200, 1650, 2210, 900], "Danh mục API thiết kế đích", 6.8)
    w.heading("Hợp đồng API trọng yếu", 2)
    for idx, contract in enumerate(API_CONTRACTS, 1):
        method, endpoint, purpose, role, scope, request, query, success, error, validation, rules, effects, idem, tests = contract
        w.heading(f"{method} {endpoint}", 3)
        w.table(
            ["Thuộc tính", "Đặc tả"],
            [["Method + Endpoint", f"{method} {endpoint}"], ["Mục đích", purpose], ["Role được phép", role],
             ["Scope dữ liệu", scope], ["Request body", request], ["Query params", query],
             ["Response success", success], ["Response error", error], ["Validation", validation],
             ["Business rules", rules], ["Side effects", effects], ["Idempotency", idem],
             ["Test case liên quan", tests]],
            [2000, 7360], f"Hợp đồng {method} {endpoint}", 7.8,
        )


def add_ui_flows(w: Writer) -> None:
    w.heading("User flow theo vai trò", 2)
    w.callout("Flow sinh viên", "Login → Dashboard → Profile Setup → Assessment → Analysis Result → Recommendation → Roadmap → Progress Update.", LIGHT_BLUE)
    w.callout("Flow cố vấn", "Login → Advisor Dashboard → Assigned Students → Student Detail → Evidence Review → Advisor Comment.", LIGHT_BLUE)
    w.callout("Flow quản trị", "Login → Admin Dashboard → Skill Catalog → Career Goal → Required Skill → Question Bank → Policy → Dashboard.", LIGHT_BLUE)
    w.heading("Từ điển màn hình và trạng thái", 2)
    w.table(
        ["Màn hình", "Role", "Mục đích", "Form/input", "Button/action", "API", "Loading/empty/error"],
        SCREEN_ROWS,
        [1200, 900, 1700, 1300, 1600, 1500, 1160],
        "Đặc tả màn hình theo vai trò",
        6.6,
    )


def add_rbac_design(w: Writer) -> None:
    w.heading("Mô hình role và scope", 2)
    w.paragraph(
        "Quyền được quyết định theo hai lớp: role xác định nhóm hành động; scope xác định tập bản ghi. "
        "Student dùng own data, Advisor/Lecturer dùng assigned students có AdvisorAssignment ACTIVE, Admin dùng all data theo quyền chi tiết."
    )
    rows = []
    for action, student, advisor, admin, control in RBAC_ROWS:
        allowed_roles = []
        if student not in {"Không", "Không dùng endpoint này"}:
            allowed_roles.append("Student")
        if advisor not in {"Không"}:
            allowed_roles.append("Advisor")
        if admin not in {"Không"}:
            allowed_roles.append("Admin")
        scope = f"Student: {student}; Advisor: {advisor}; Admin: {admin}"
        audit = "Có" if action.startswith(("POST", "PATCH", "DELETE")) or "activity-logs" in action or "policies" in action else "Không; chỉ access log"
        rows.append([action, ", ".join(allowed_roles) or "Hệ thống", scope, control, audit, "403 hoặc 404 an toàn"])
    w.table(
        ["API/action", "Role", "Dữ liệu được thao tác", "Kiểm tra owner/assignment", "ActivityLog", "Khi từ chối"],
        rows,
        [1800, 900, 2450, 1900, 1200, 1110],
        "Ma trận role-scope và kiểm soát bắt buộc",
        6.6,
    )
    w.heading("Nguyên tắc kiểm soát ngang", 2)
    w.bullets([
        "Không tin studentId từ body ở luồng tự phục vụ; subject luôn lấy từ access token.",
        "Cố vấn chỉ đọc hoặc ghi dữ liệu của sinh viên có AdvisorAssignment ACTIVE tại thời điểm truy vấn.",
        "Admin thao tác dữ liệu nhạy cảm phải ghi ActivityLog với lý do và requestId.",
        "Dashboard, list, count và export áp dụng scope ngay trong where/join, không lọc sau khi tải toàn bộ.",
        "Ưu tiên 404 khi cần tránh tiết lộ sự tồn tại của bản ghi ngoài phạm vi.",
    ])


def test_api_uc(code: str) -> str:
    prefix = code.split("-")[1]
    return {
        "AUTH": "/api/auth/* / UC-ACC-01..02", "USR": "/api/users/* / UC-ACC-03",
        "ADV": "/api/advisor-* / UC-ACC-04,UC-ADV-01", "CAT": "/api/skills|career-goals / UC-CAT-01..02",
        "PRO": "/api/profiles/* / UC-PRO-01..02", "EVD": "/api/evidences/* / UC-EVD-01..02",
        "ASM": "/api/assessments|questions / UC-ASM-01..04", "ENG": "/api/competency-analysis/* / UC-ENG-01..02",
        "REC": "/api/recommendations|learning-resources / UC-REC-01..03", "ROA": "/api/roadmaps/* / UC-ROA-01..02",
        "DSH": "/api/dashboard/* / UC-DSH-01", "NOT": "/api/notifications/* / UC-NOT-01..02",
        "EXP": "/api/exports/* / UC-EXP-01..02", "SEC": "Nhiều API / kiểm thử bảo mật",
    }.get(prefix, "Theo module / use case liên quan")


def test_steps(module: str) -> str:
    if module in {"Engine", "Assessment"}:
        return "1) Seed context/policy; 2) Gọi API; 3) Kiểm tra response; 4) Đối chiếu transaction, snapshot và bất biến."
    if module in {"Security", "Evidence", "Export"}:
        return "1) Tạo actor và scope; 2) Gửi input hợp lệ/độc hại; 3) Kiểm tra status; 4) Kiểm tra không rò dữ liệu, file và audit."
    return "1) Chuẩn bị dữ liệu; 2) Gọi API/action; 3) Kiểm tra response; 4) Kiểm tra DB và side effect."


def add_test_design(w: Writer) -> None:
    w.heading("Chiến lược kiểm thử", 2)
    w.paragraph(
        "Bộ test kết hợp unit cho quy tắc, property test cho bất biến Engine, integration cho transaction/database, "
        "contract cho request-response, E2E theo role và security test cho ownership/scope."
    )
    rows = []
    for code, module, condition_input, expected in TEST_CASES:
        condition = condition_input.split(";", 1)[0]
        rows.append([code, module, condition, condition_input, test_steps(module), expected, test_api_uc(code)])
    w.table(
        ["Mã TC", "Module", "Điều kiện", "Input", "Các bước", "Kết quả mong đợi", "API/use case"],
        rows,
        [850, 800, 1200, 1350, 2100, 1900, 1160],
        "Test case chức năng, tích hợp và bảo mật",
        6.3,
    )
    w.heading("Tiêu chí nghiệm thu và phát hành", 2)
    w.bullets([
        "100% test P0/P1 đạt; không còn lỗi phân quyền mức cao hoặc rất cao.",
        "Mọi yêu cầu có use case, API, dữ liệu, màn hình và test tương ứng trong ma trận truy vết.",
        "Migration thử nghiệm trên bản sao dữ liệu; rollback hoặc forward-fix được diễn tập.",
        "Contract frontend/backend được kiểm tra từ OpenAPI; enum và state machine thống nhất.",
        "Hiệu năng đạt NFR cho Engine, dashboard, list và export job; upload/export kiểm tra hạn và ownership.",
    ])


def trace_rule(fr: str) -> str:
    prefix = fr.split("-")[1]
    return {
        "ACC": "BR-ACC", "CAT": "BR-CAT", "PRO": "BR-PRO", "EVD": "BR-EVD", "ASM": "BR-ASM",
        "ENG": "BR-ENG", "REC": "BR-REC", "ROA": "BR-ROA", "ADV": "BR-ADV", "DSH": "BR-DSH",
        "NOT": "BR-NOT", "EXP": "BR-EXP", "AUD": "BR-AUD",
    }.get(prefix, "BR-GEN")


def trace_screen(fr: str) -> str:
    prefix = fr.split("-")[1]
    return {
        "ACC": "Đăng nhập/Tài khoản", "CAT": "Danh mục/Mục tiêu", "PRO": "Hồ sơ", "EVD": "Minh chứng",
        "ASM": "Assessment/Lịch sử", "ENG": "Kết quả phân tích", "REC": "Gợi ý/Tài nguyên", "ROA": "Lộ trình",
        "ADV": "Sinh viên được phân công/Nhận xét", "DSH": "Dashboard theo role", "NOT": "Thông báo",
        "EXP": "Portfolio/Báo cáo", "AUD": "Audit log",
    }.get(prefix, "Theo module")


def add_traceability(w: Writer) -> None:
    w.paragraph(
        "Ma trận nối yêu cầu với use case, business rule, API, bảng dữ liệu, màn hình và test. "
        "Cột ghi chú chỉ mô tả quyết định thiết kế, không phản ánh trạng thái triển khai."
    )
    rows = []
    for fr, uc, api, data, test, _status in TRACEABILITY:
        rows.append([fr, uc, trace_rule(fr), api, data, trace_screen(fr), test, "Thiết kế đích; nghiệm thu theo Chương 14"])
    w.table(
        ["FR", "Use case", "Business rule", "API", "Bảng dữ liệu", "Màn hình", "Test case", "Ghi chú thiết kế"],
        rows,
        [850, 850, 850, 1700, 1500, 1200, 1100, 1310],
        "Ma trận truy vết FR - UC - BR - API - DB - UI - Test",
        6.0,
    )
    groups = {}
    for row in rows:
        fr = row[0]
        prefix = fr.split("-")[1]
        count = 1
        m = re.search(r"(\d+)\.\.(\d+)", fr)
        if m:
            count = int(m.group(2)) - int(m.group(1)) + 1
        groups[prefix] = groups.get(prefix, 0) + count
    summary = []
    for prefix, count in sorted(groups.items()):
        summary.append([prefix, str(count), str(count), str(count), str(count), str(count), str(count), "Không; còn quyết định chính sách được liệt kê ở Chương 16"])
    w.heading("Traceability Coverage Summary", 2)
    w.table(
        ["Nhóm chức năng", "Số FR", "Có use case", "Có API", "Có DB", "Có UI", "Có test", "Gap còn lại"],
        summary,
        [1300, 700, 900, 800, 750, 750, 800, 3360],
        "Tổng hợp độ phủ truy vết",
        6.8,
    )


def add_state_machines(w: Writer) -> None:
    w.heading("Máy trạng thái nghiệp vụ", 2)
    w.paragraph(
        "State transition phải được kiểm tra tại service trong cùng transaction với thay đổi dữ liệu. "
        "Client chỉ hiển thị hành động được phép, không phải nơi quyết định quyền chuyển trạng thái."
    )
    for entity, rows in STATE_MACHINES.items():
        w.heading(entity, 3)
        w.table(
            ["Trạng thái", "Ý nghĩa", "Được chuyển tới", "Actor/Hệ thống", "Điều kiện", "Side effect"],
            rows,
            [1150, 1600, 1350, 1250, 2100, 1910],
            f"Máy trạng thái {entity}",
            6.8,
        )


def add_open_questions(w: Writer) -> None:
    w.heading("Câu hỏi cần chốt trước triển khai", 2)
    w.numbered([
        "AssessmentPolicy và AnalysisPolicy được lưu có version trong database hay đóng gói theo release?",
        "Tài khoản được đăng ký công khai hay chỉ do nhà trường tạo/import?",
        "Loại minh chứng nào có thời hạn, ai được xác minh và thời gian xử lý tối đa là bao lâu?",
        "Mức tự khai được dùng làm currentLevel hay chỉ là dữ liệu tham khảo có confidence thấp?",
        "Lịch sử Assessment có hiển thị đáp án đúng và lời giải ngay sau khi nộp không?",
        "Ba lần đánh giá được tính theo student+context, học kỳ hay policyVersion?",
        "Feedback gợi ý chỉ phục vụ báo cáo hay được dùng để điều chỉnh ranking/template?",
        "Roadmap dependency HARD và SOFT được áp dụng cho nhóm mục nào?",
        "Portfolio PDF mặc định công khai trường nào và có QR/link xác thực không?",
        "Báo cáo XLSX tối đa bao nhiêu dòng trước khi bắt buộc chia file?",
        "Retention của ActivityLog, EvidenceRevision và file export là bao lâu?",
        "Email notification nằm trong MVP hay chỉ dùng kênh IN_APP?",
    ])


def remove_table_column(table: Table, index: int) -> None:
    grid = table._tbl.tblGrid
    grid_cols = list(grid)
    if 0 <= index < len(grid_cols):
        grid.remove(grid_cols[index])
    for row in table.rows:
        cells = row.cells
        if 0 <= index < len(cells):
            row._tr.remove(cells[index]._tc)


def remove_banned_columns(doc: Document) -> None:
    banned = {
        "hiện trạng", "hiện trạng code", "trạng thái triển khai", "đã có trong code", "đã triển khai",
        "thiếu trong source", "file/service/controller liên quan",
    }
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [re.sub(r"\s+", " ", c.text.strip().lower()) for c in table.rows[0].cells]
        indices = [i for i, h in enumerate(headers) if h in banned]
        for idx in reversed(indices):
            remove_table_column(table, idx)
        if indices:
            widths = [TABLE_WIDTH // len(table.columns)] * len(table.columns)
            widths[-1] += TABLE_WIDTH - sum(widths)
            format_table(table, widths, 7.5)


def strip_number(text: str) -> str:
    text = re.sub(r"^\d+(?:\.\d+)*\.\s*", "", text.strip())
    return text


def renumber_headings(doc: Document) -> None:
    chapter = h2 = h3 = 0
    for p in doc.paragraphs:
        style = p.style.name if p.style else ""
        text = p.text.strip()
        if style == "Heading 1":
            if text in {"Thông tin tài liệu", "Mục lục"}:
                continue
            chapter += 1
            h2 = h3 = 0
            p.text = f"{chapter}. {strip_number(text)}"
        elif style == "Heading 2" and chapter:
            h2 += 1
            h3 = 0
            p.text = f"{chapter}.{h2}. {strip_number(text)}"
        elif style == "Heading 3" and chapter and h2:
            h3 += 1
            p.text = f"{chapter}.{h2}.{h3}. {strip_number(text)}"
        if style.startswith("Heading"):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True


def renumber_captions(doc: Document) -> None:
    counts = {"Bảng": 0, "Hình": 0}
    for p in doc.paragraphs:
        if not p.style or p.style.name != "Caption":
            continue
        m = re.match(r"^(Bảng|Hình)\s+\d+\.\s*(.*)$", p.text.strip())
        if not m:
            continue
        label, caption = m.groups()
        counts[label] += 1
        p.text = f"{label} {counts[label]}. {caption}"
        p.paragraph_format.keep_with_next = True


def sanitize_text(doc: Document) -> None:
    replacements = {
        "Định hướng: Đặc tả triển khai toàn hệ thống và đối chiếu mã nguồn": "Định hướng: Đặc tả thiết kế đích toàn hệ thống trước khi lập trình",
        "đối chiếu code và nghiệm thu": "truy vết yêu cầu và nghiệm thu",
        "truy vết và đối chiếu hiện trạng triển khai": "truy vết và xác nhận tính nhất quán thiết kế",
        "Tài liệu gồm hai lớp: thiết kế đích đủ chi tiết để phát triển và phần đối chiếu với hiện trạng mã nguồn.": "Tài liệu trình bày một mạch thiết kế đích đủ chi tiết để phát triển và kiểm thử.",
        "API dự kiến/hiện có": "API liên quan",
        " (đề xuất)": "",
        "GET /api/assessments/history": "GET /api/assessments/me",
        "POST /api/exports/reports": "POST /api/exports/report",
    }
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for p in paragraphs:
        for run in p.runs:
            original = run.text
            text = original
            for old, new in replacements.items():
                text = text.replace(old, new)
            if text != original:
                run.text = text


def update_front_matter(doc: Document) -> None:
    for p in doc.paragraphs:
        if p.text.strip() == "Phiên bản tài liệu: 3.0":
            p.text = "Phiên bản tài liệu: 4.0 - Thiết kế đích thống nhất"
        elif p.text.strip().startswith("Định hướng:"):
            p.text = "Định hướng: Đặc tả thiết kế đích toàn hệ thống trước khi lập trình"
    if doc.tables:
        for row in doc.tables[0].rows:
            key = row.cells[0].text.strip()
            if key == "Mục đích":
                row.cells[1].text = "Làm cơ sở phê duyệt thiết kế, phát triển backend/frontend, kiểm thử và nghiệm thu toàn hệ thống."
            elif key == "Phương pháp":
                row.cells[1].text = "Phân tích chức năng, use case, mô hình dữ liệu, hợp đồng API, đặc tả UI, bảo mật, kiểm thử và truy vết yêu cầu."
            elif key == "Trạng thái":
                row.cells[1].text = "Bản thiết kế đích hoàn chỉnh để rà soát và phê duyệt trước triển khai."
        format_table(doc.tables[0], [2300, 7060], 8.5)


def validate_design_only(doc: Document) -> None:
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            texts.extend(c.text for c in row.cells)
    all_text = "\n".join(texts).lower()
    banned = [
        "đối chiếu với hiện trạng mã nguồn", "danh mục api hiện có", "hiện trạng code",
        "đã có trong code", "đã triển khai", "thiếu trong source", "file/service/controller liên quan",
    ]
    present = [x for x in banned if x in all_text]
    if present:
        raise RuntimeError(f"Còn nội dung bị cấm: {present}")
    h1 = [p.text for p in doc.paragraphs if p.style and p.style.name == "Heading 1" and p.text not in {"Thông tin tài liệu", "Mục lục"}]
    if len(h1) != 16:
        raise RuntimeError(f"Số chương không đúng: {len(h1)} - {h1}")


def build() -> None:
    doc = Document(DOCX)
    source_elements = [deepcopy(el) for el in list(doc.element.body)]
    body = doc.element.body
    sect_pr = body.sectPr
    for el in list(body):
        if el is not sect_pr:
            body.remove(el)

    for el in source_elements[:39]:
        append_element(doc, el)
    add_toc(doc)
    w = Writer(doc)

    w.heading(CHAPTERS[0], 1)
    append_range(doc, source_elements, 43, 57)

    w.heading(CHAPTERS[1], 1, new_page=True)
    append_range(doc, source_elements, 57, 81)

    w.heading(CHAPTERS[2], 1, new_page=True)
    append_range(doc, source_elements, 82, 127)

    w.heading(CHAPTERS[3], 1, new_page=True)
    append_range(doc, source_elements, 572, 577)
    add_function_design(w)

    w.heading(CHAPTERS[4], 1, new_page=True)
    append_range(doc, source_elements, 578, 694)
    add_missing_use_cases(w)

    w.heading(CHAPTERS[5], 1, new_page=True)
    append_range(doc, source_elements, 216, 237)

    w.heading(CHAPTERS[6], 1, new_page=True)
    append_range(doc, source_elements, 366, 393)

    w.heading(CHAPTERS[7], 1, new_page=True)
    append_range(doc, source_elements, 394, 420)
    append_range(doc, source_elements, 733, 906)
    add_missing_entities(w)

    w.heading(CHAPTERS[8], 1, new_page=True)
    append_range(doc, source_elements, 695, 732)
    add_question_coverage(w)

    w.heading(CHAPTERS[9], 1, new_page=True)
    add_skill_scale(w)
    append_range(doc, source_elements, 238, 365)

    w.heading(CHAPTERS[10], 1, new_page=True)
    append_range(doc, source_elements, 476, 492)
    add_api_design(w)

    w.heading(CHAPTERS[11], 1, new_page=True)
    add_ui_flows(w)
    append_range(doc, source_elements, 973, 997)

    w.heading(CHAPTERS[12], 1, new_page=True)
    append_range(doc, source_elements, 452, 475)
    add_rbac_design(w)

    w.heading(CHAPTERS[13], 1, new_page=True)
    add_test_design(w)
    append_range(doc, source_elements, 493, 515)

    w.heading(CHAPTERS[14], 1, new_page=True)
    add_traceability(w)

    w.heading(CHAPTERS[15], 1, new_page=True)
    w.heading("Danh mục quyết định thiết kế", 2)
    append_range(doc, source_elements, 542, 546)
    add_state_machines(w)
    w.heading("Bảng thuật ngữ viết tắt", 2)
    append_range(doc, source_elements, 559, 571)
    add_open_questions(w)
    w.callout(
        "Kết thúc tài liệu",
        "Phiên bản 4.0 là tài liệu thiết kế đích thống nhất, nối chức năng → use case → quy tắc → API → dữ liệu → giao diện → kiểm thử. "
        "Competency Analysis Engine giữ vai trò lõi nhưng được đặt trong đầy đủ luồng đầu vào, đầu ra và vận hành toàn hệ thống.",
        GREEN,
    )

    enhance_use_case_tables(doc)
    insert_entity_contexts(doc)
    remove_banned_columns(doc)
    sanitize_text(doc)
    renumber_headings(doc)
    renumber_captions(doc)
    update_front_matter(doc)
    validate_design_only(doc)

    doc.core_properties.title = "Phân tích và thiết kế hệ thống Student Digital Competency Profile - Thiết kế đích"
    doc.core_properties.subject = "Chức năng, use case, dữ liệu, Engine, API, UI, bảo mật, kiểm thử và truy vết"
    doc.core_properties.comments = "Phiên bản 4.0 - 29/06/2026"
    doc.save(TMP)
    os.replace(TMP, DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
