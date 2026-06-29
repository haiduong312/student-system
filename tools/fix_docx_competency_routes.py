from pathlib import Path

from docx import Document


DOCX = Path(
    "/Users/nguyenhaiduong/Downloads/Student-system/"
    "PTTK_Student_Digital_Competency_Profile_Ban_Chuyen_Nghiep_Toan_He_Thong.docx"
)


def replace_in_paragraph(paragraph, old: str, new: str) -> int:
    count = 0
    for run in paragraph.runs:
        if old in run.text:
            count += run.text.count(old)
            run.text = run.text.replace(old, new)
    return count


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def main() -> None:
    document = Document(DOCX)
    replacements = 0
    for paragraph in document.paragraphs:
        replacements += replace_in_paragraph(
            paragraph, "/competency/", "/competency-analysis/"
        )
    for table in document.tables:
        for paragraph in iter_table_paragraphs(table):
            replacements += replace_in_paragraph(
                paragraph, "/competency/", "/competency-analysis/"
            )
    document.save(DOCX)
    print(f"Đã sửa {replacements} route trong {DOCX.name}")


if __name__ == "__main__":
    main()
