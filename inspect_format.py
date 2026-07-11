from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

src=Path('/Users/nguyenhaiduong/Downloads/PTTK_Student_Digital_Competency_Profile_Da_Chinh_Dinh_Dang_Muc_9.5.docx')
d=Document(src)
print('sections', len(d.sections), [(s.page_width, s.page_height, s.top_margin, s.bottom_margin) for s in d.sections])
for i,p in enumerate(d.paragraphs):
    if p.text.startswith('5.4.') or p.text.startswith('Bảng 14'):
        print('P',i,p.text,p.style.name,p.paragraph_format.keep_with_next,p.paragraph_format.page_break_before)
        for r in p.runs:
            print(' run',repr(r.text),r.font.name,r.font.size,r.bold,r.italic)
for ti,t in enumerate(d.tables):
    if len(t.rows)>1 and t.cell(1,0).text.strip()=='Mã use case':
        print('table',ti,t.style.name,len(t.rows),len(t.columns))
        print('widths',[c.width for c in t.rows[0].cells])
        for ri in [0,1,8,9]:
            row=t.rows[ri]
            print(' row',ri,'cant',row._tr.find(qn('w:trPr')))
            for ci,c in enumerate(row.cells):
                p=c.paragraphs[0]
                print(ci,repr(c.text[:100]),p.style.name,[(r.font.name,r.font.size.pt if r.font.size else None,r.bold) for r in p.runs[:2]])
        break
print('settings updateFields',d.settings.element.find(qn('w:updateFields')) is not None)
